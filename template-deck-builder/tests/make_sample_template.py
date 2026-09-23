#!/usr/bin/env python3
"""Generate a sample corporate template for testing: 16:9, custom theme colors and
fonts, renamed and custom layouts (Cover, Title and Content, Section Divider, Two Column,
Title and Chart, Quote, Title Only, Blank), a dedicated Source placeholder on content
layouts, and a decorative accent bar on the master. Writes .pptx and .potx.

Usage: python3 make_sample_template.py OUT_DIR
"""
import copy
import io
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}

COLORS = {"dk1": "1A1A1A", "lt1": "FFFFFF", "dk2": "14213D", "lt2": "E7E9EE",
          "accent1": "1F5AA6", "accent2": "D08C2B", "accent3": "5B6B82", "accent4": "9FB3CF",
          "accent5": "B5462F", "accent6": "2B8A9E", "hlink": "1F5AA6", "folHlink": "5B6B82"}
FONTS = {"major": "Georgia", "minor": "Arial"}


def set_theme(prs):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    root.set("name", "Sample Corporate")
    scheme = root.find(".//a:clrScheme", NS)
    scheme.set("name", "Sample Corporate")
    for child in list(scheme):
        tag = etree.QName(child).localname
        for c in list(child):
            child.remove(c)
        etree.SubElement(child, f"{{{A}}}srgbClr", val=COLORS[tag])
    for kind in ("major", "minor"):
        root.find(f".//a:fontScheme/a:{kind}Font/a:latin", NS).set("typeface", FONTS[kind])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def scale_x(shapes, k):
    """Only shapes with their OWN xfrm. Setting .left on a placeholder that inherits
    its position creates a new xfrm with y=0 and height=0 (python-pptx gotcha)."""
    for s in shapes:
        if s._element.find("p:spPr/a:xfrm", NS) is not None:
            s.left, s.width = Emu(int(s.left * k)), Emu(int(s.width * k))


def style_master(prs):
    m = prs.slide_master
    tx = m._element.find(".//p:txStyles", NS)
    t1 = tx.find("p:titleStyle/a:lvl1pPr", NS)
    t1.set("algn", "l")
    t1.find("a:defRPr", NS).set("sz", "2600")
    t1.find("a:defRPr", NS).set("b", "0")
    for lvl, sz in (("lvl1pPr", "1800"), ("lvl2pPr", "1600"), ("lvl3pPr", "1400")):
        tx.find(f"p:bodyStyle/a:{lvl}", NS).find("a:defRPr", NS).set("sz", sz)
    # decorative accent bar on the master (a template shape, not content)
    bar = etree.fromstring(
        f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="90" name="Accent Bar"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{Inches(0.6)}" y="{Inches(0.25)}"/><a:ext cx="{Inches(0.9)}" cy="{Inches(0.06)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="accent2"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
    m.shapes._spTree.append(bar)


def add_placeholder(layout, clone_from, idx, name, box, size_pt=None, bullets=True):
    el = copy.deepcopy(clone_from._element)
    el.find(".//p:cNvPr", NS).set("id", str(100 + idx))
    el.find(".//p:cNvPr", NS).set("name", name)
    ph = el.find(".//p:nvPr/p:ph", NS)
    ph.set("idx", str(idx))
    ph.attrib.pop("type", None)
    ph.set("type", "body")
    ph.set("sz", "quarter")
    sppr = el.find("p:spPr", NS)
    for old in sppr.findall("a:xfrm", NS):
        sppr.remove(old)
    xfrm = etree.Element(f"{{{A}}}xfrm")
    etree.SubElement(xfrm, f"{{{A}}}off", x=str(int(box[0])), y=str(int(box[1])))
    etree.SubElement(xfrm, f"{{{A}}}ext", cx=str(int(box[2])), cy=str(int(box[3])))
    sppr.insert(0, xfrm)
    txb = el.find(".//p:txBody", NS)
    for old in txb.findall("a:lstStyle", NS):
        txb.remove(old)
    lst = etree.SubElement(txb, f"{{{A}}}lstStyle")
    txb.insert(1, lst)
    lvl = etree.SubElement(lst, f"{{{A}}}lvl1pPr", marL="0", indent="0")
    if not bullets:
        etree.SubElement(lvl, f"{{{A}}}buNone")
    if size_pt:
        etree.SubElement(lvl, f"{{{A}}}defRPr", sz=str(int(size_pt * 100)))
    for p in txb.findall("a:p", NS):
        txb.remove(p)
    p = etree.SubElement(txb, f"{{{A}}}p")
    r = etree.SubElement(p, f"{{{A}}}r")
    etree.SubElement(r, f"{{{A}}}rPr", lang="en-US")
    etree.SubElement(r, f"{{{A}}}t").text = name
    layout.shapes._spTree.append(el)


def reset_text_style(ph):
    """Empty the placeholder's lstStyle and bodyPr so it inherits from the master."""
    txb = ph._element.find(".//p:txBody", NS)
    for tag in ("a:bodyPr", "a:lstStyle"):
        el = txb.find(tag, NS)
        for c in list(el):
            el.remove(c)
        el.attrib.clear()


def ph_by_type(layout, t):
    return [p for p in layout.placeholders if p.placeholder_format.type == t]


def main(out_dir):
    from pptx.enum.shapes import PP_PLACEHOLDER as T
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    k = (13.333 / 10.0)
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    scale_x(prs.slide_master.shapes, k)
    for l in prs.slide_layouts:
        scale_x(l.shapes, k)
    set_theme(prs)
    style_master(prs)

    L = list(prs.slide_layouts)
    cover_title = ph_by_type(L[0], T.CENTER_TITLE)[0]
    lst = cover_title._element.find(".//a:lstStyle", NS)
    lvl = etree.SubElement(lst, f"{{{A}}}lvl1pPr", algn="ctr")
    etree.SubElement(lvl, f"{{{A}}}defRPr", sz="3600")
    names = {0: "Cover", 1: "Title and Content", 2: "Section Divider", 3: "Two Column",
             5: "Title Only", 6: "Blank", 7: "Quote", 8: "Title and Chart"}
    for i, n in names.items():
        L[i].name = n
    src_box = (Inches(0.6), Inches(6.88), Inches(12.1), Inches(0.32))

    # content layouts: tighten geometry to a consulting grid, make every title and body
    # inherit the master styles (the stock layouts override sizes), add a Source line
    for i in (1, 3, 8):
        lay = L[i]
        title = ph_by_type(lay, T.TITLE)[0]
        title.left, title.top, title.width, title.height = Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.1)
        for ph in lay.placeholders:
            if ph.placeholder_format.type in (T.TITLE, T.OBJECT, T.BODY, T.PICTURE):
                reset_text_style(ph)
        body = [p for p in lay.placeholders if p.placeholder_format.type in (T.OBJECT, T.BODY, T.PICTURE)]
        add_placeholder(lay, body[0], 20, "Source", src_box, size_pt=10, bullets=False)
    content = ph_by_type(L[1], T.OBJECT)[0]
    content.left, content.top, content.width, content.height = Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.0)
    two = ph_by_type(L[3], T.OBJECT)
    for j, p in enumerate(sorted(two, key=lambda p: p.left)):
        p.left, p.top, p.width, p.height = Inches(0.6 + j * 6.2), Inches(1.7), Inches(5.9), Inches(5.0)

    # Title and Chart: picture placeholder becomes a chart placeholder; drop the caption
    lay = L[8]
    pic = ph_by_type(lay, T.PICTURE)[0]
    pic._element.find(".//p:nvPr/p:ph", NS).set("type", "chart")
    pic.left, pic.top, pic.width, pic.height = Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.0)
    for p in ph_by_type(lay, T.BODY):
        if p.placeholder_format.idx != 20:
            p._element.getparent().remove(p._element)

    # Quote: keep one body as the quote, add an attribution line, drop title + content
    lay = L[7]
    for p in list(lay.placeholders):
        if p.placeholder_format.type in (T.TITLE, T.OBJECT):
            p._element.getparent().remove(p._element)
    q = ph_by_type(lay, T.BODY)[0]
    q.left, q.top, q.width, q.height = Inches(1.6), Inches(1.9), Inches(10.1), Inches(2.9)
    q.name = "Quote"
    add_placeholder(lay, q, 21, "Attribution", (Inches(1.6), Inches(5.0), Inches(10.1), Inches(0.6)), size_pt=14, bullets=False)
    txb = q._element.find(".//p:txBody", NS)
    lst = txb.find("a:lstStyle", NS)
    if lst is None:
        lst = etree.Element(f"{{{A}}}lstStyle")
        txb.insert(1, lst)
    for old in lst.findall("a:lvl1pPr", NS):
        lst.remove(old)
    lvl = etree.Element(f"{{{A}}}lvl1pPr", marL="0", indent="0")
    lst.insert(0, lvl)
    etree.SubElement(lvl, f"{{{A}}}buNone")
    etree.SubElement(lvl, f"{{{A}}}defRPr", sz="2800", i="1")

    # drop layouts the house style does not use (Comparison, vertical text ones)
    for i in (10, 9, 4):
        prs.slide_layouts.remove(L[i])

    pptx_path = out_dir / "sample_template.pptx"
    prs.save(pptx_path)
    # .potx copy: identical package, template main content type
    data = io.BytesIO()
    zin = zipfile.ZipFile(pptx_path)
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            b = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                b = b.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            zout.writestr(item, b)
    (out_dir / "sample_template.potx").write_bytes(data.getvalue())
    print(f"Wrote {pptx_path} and {out_dir / 'sample_template.potx'}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
