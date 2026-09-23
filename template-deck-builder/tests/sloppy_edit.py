#!/usr/bin/env python3
"""Simulate a hand-edited deck that ignored the template: an off-theme font and
color on slide 1, plus an extra slide 'remade from scratch' on the Blank layout
with drawn boxes. Used to prove qa_deck.py catches template bypass.
Usage: python3 sloppy_edit.py IN.pptx OUT.pptx"""
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

prs = Presentation(sys.argv[1])
title = prs.slides[0].shapes.title
for p in title.text_frame.paragraphs:
    for r in p.runs:
        r.font.name = "Comic Sans MS"
        r.font.color.rgb = RGBColor(0xFF, 0x00, 0xFF)
blank = next(l for l in prs.slide_layouts if l.name == "Blank")
s = prs.slides.add_slide(blank)
box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(12.3), Inches(1.2))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x33, 0x99, 0xFF)
tb = s.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(11), Inches(0.8))
tb.text_frame.text = "Next steps"
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
tb.text_frame.paragraphs[0].runs[0].font.name = "Helvetica Neue"
prs.save(sys.argv[2])
print("Wrote", sys.argv[2])
