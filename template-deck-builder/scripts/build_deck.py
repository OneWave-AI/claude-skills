#!/usr/bin/env python3
"""Build (or edit) a deck from a storyboard JSON by filling the TEMPLATE's own
layouts and placeholders. Never draws free-floating text boxes unless the spec
asks for one (`free_text`), or a layout has no source placeholder and a
`source` line must go somewhere (named 'tdb-source', reported in the log).

Usage:
  python3 build_deck.py storyboard.json --template T.pptx -o out.pptx
  python3 build_deck.py edits.json --edit existing.pptx -o out.pptx

See SKILL.md for the storyboard schema.
"""
import argparse
import copy
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION  # noqa: E402
from pptx.enum.dml import MSO_THEME_COLOR  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.shapes.placeholder import ChartPlaceholder, TablePlaceholder  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from inspect_template import classify  # noqa: E402
from tdb_common import (NS, PP_PLACEHOLDER, TITLE_TYPES, content_placeholders, find_layout,  # noqa: E402
                        is_source_ph, open_presentation, ph_type, rect, title_placeholder)

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_column_100": XL_CHART_TYPE.COLUMN_STACKED_100, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS,
    "line_plain": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
}
ACCENTS = [MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.ACCENT_2, MSO_THEME_COLOR.ACCENT_3,
           MSO_THEME_COLOR.ACCENT_4, MSO_THEME_COLOR.ACCENT_5, MSO_THEME_COLOR.ACCENT_6]
NUMERIC_RE = re.compile(r"^[\s$€£+\-(]*[\d.,]+\s*(%|x|pp|pts?|bps|[kKmMbB]n?)?\)?\s*$")
CHART_OK = {PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.OBJECT}
TABLE_OK = {PP_PLACEHOLDER.TABLE, PP_PLACEHOLDER.OBJECT}


class Log:
    def __init__(self):
        self.slides, self.errors, self.warnings = [], [], []

    def err(self, n, msg):
        self.errors.append(f"slide {n}: {msg}")

    def warn(self, n, msg):
        self.warnings.append(f"slide {n}: {msg}")


# ---------------------------------------------------------------- text

def _runs_from_markup(p, text, base_rpr=None):
    """Write `text` into paragraph p; **x** becomes a bold run."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        r = p.add_run()
        if base_rpr is not None:
            old = r._r.find("a:rPr", NS)
            if old is not None:
                r._r.remove(old)
            r._r.insert(0, copy.deepcopy(base_rpr))
        if part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]
            r.font.bold = True
        else:
            r.text = part


def _flatten(items, level=0):
    """str | list (nested lists = sub-bullets) | {'text':..,'level':..} -> [(level, text)]."""
    if items is None:
        return []
    if isinstance(items, str):
        return [(level, line) for line in items.split("\n")]
    if isinstance(items, dict):
        return [(int(items.get("level", level)), items["text"])]
    out = []
    for it in items:
        if isinstance(it, list):
            out.extend(_flatten(it, level + 1))
        else:
            out.extend(_flatten(it, level))
    return out


def write_text(shape, items, preserve=True):
    """Replace the text of a placeholder, keeping the first run's formatting
    (so edits to an existing deck keep its look) and the inherited styles."""
    tf = shape.text_frame
    base_rpr = None
    if preserve:
        r0 = tf._txBody.find(".//a:p/a:r/a:rPr", NS)
        base_rpr = copy.deepcopy(r0) if r0 is not None else None
        if base_rpr is not None:
            base_rpr.attrib.pop("b", None)
    lines = _flatten(items)
    tf.clear()
    for i, (lvl, text) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for r in list(p.runs):
            p._p.remove(r._r)
        p.level = max(0, min(lvl, 8))
        _runs_from_markup(p, text, base_rpr)


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


# ---------------------------------------------------------------- charts

def insert_chart(slide, ph, spec, n, log):
    kind = spec.get("type", "column")
    if kind not in CHART_TYPES:
        log.err(n, f"chart type '{kind}' unknown; use one of {sorted(CHART_TYPES)}")
        return None
    cd = CategoryChartData(number_format=spec.get("number_format", "General"))
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["name"], s["values"])
    ct = CHART_TYPES[kind]
    if ph_type(ph) in CHART_OK:
        gf = ChartPlaceholder(ph._element, ph._parent).insert_chart(ct, cd)
    else:  # text-only body placeholder: take its geometry, then drop it
        l, t, w, h = rect(ph)
        gf = slide.shapes.add_chart(ct, Emu(l), Emu(t), Emu(w), Emu(h), cd)
        gf.name = f"tdb-chart@ph{ph.placeholder_format.idx}"
        remove_shape(ph)
        log.warn(n, "chart placed in the geometry of a text-only BODY placeholder (template has no chart/content placeholder here)")
    style_chart(gf.chart, spec, kind)
    return gf


def style_chart(chart, spec, kind):
    """Theme colors only (schemeClr), no gridlines, direct labels: consulting-clean."""
    chart.font.size = Pt(spec.get("font_size", 12))
    plot = chart.plots[0]
    multi = len(spec["series"]) > 1
    is_pie = kind in ("pie", "doughnut")
    chart.has_legend = multi or is_pie
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM if not is_pie else XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
    if "title" in spec:
        chart.has_title = True
        chart.chart_title.text_frame.text = spec["title"]
    else:
        chart.has_title = False
    labels = spec.get("data_labels", True)
    plot.has_data_labels = labels
    if labels:
        dl = plot.data_labels
        dl.number_format = spec.get("label_format", spec.get("number_format", "General"))
        dl.number_format_is_linked = False
        dl.font.size = Pt(spec.get("font_size", 12))
        if kind in ("column", "bar"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    if kind.startswith(("column", "bar", "stacked")):
        plot.gap_width = spec.get("gap_width", 60)
        if kind.startswith("stacked"):
            plot.overlap = 100
    if not is_pie:
        va = chart.value_axis
        va.has_major_gridlines = False
        va.tick_labels.number_format = spec.get("number_format", "General")
        va.tick_labels.number_format_is_linked = False
        if labels and spec.get("hide_value_axis", True):
            va.visible = False
        if kind.startswith(("bar", "stacked_bar")):
            chart.category_axis.reverse_order = True  # first category on top, as read
        if "value_axis_title" in spec:
            va.has_title = True
            va.axis_title.text_frame.text = spec["value_axis_title"]
    highlight = spec.get("highlight")
    if highlight is not None and not isinstance(highlight, list):
        highlight = [highlight]
    for i, s in enumerate(plot.series):
        color = ACCENTS[i % len(ACCENTS)]
        if kind.startswith("line"):
            s.format.line.color.theme_color = color
            s.format.line.width = Pt(2.25)
            s.smooth = False
            continue
        if is_pie:
            for j in range(len(spec["categories"])):
                pt = s.points[j]
                pt.format.fill.solid()
                pt.format.fill.fore_color.theme_color = ACCENTS[j % len(ACCENTS)]
            continue
        s.format.fill.solid()
        s.format.fill.fore_color.theme_color = color
        if highlight is not None and not multi:
            for j, cat in enumerate(spec["categories"]):
                pt = s.points[j]
                pt.format.fill.solid()
                if j in highlight or cat in highlight:
                    pt.format.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
                else:  # muted theme neutral so the one message pops
                    pt.format.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_2
                    pt.format.fill.fore_color.brightness = -0.25


# ---------------------------------------------------------------- tables

def insert_table(slide, ph, spec, n, log):
    header = spec.get("columns")
    rows = spec["rows"]
    data = ([header] if header else []) + rows
    nr, nc = len(data), max(len(r) for r in data)
    if ph_type(ph) in TABLE_OK:
        gf = TablePlaceholder(ph._element, ph._parent).insert_table(nr, nc)
    else:
        l, t, w, h = rect(ph)
        gf = slide.shapes.add_table(nr, nc, Emu(l), Emu(t), Emu(w), Emu(h))
        gf.name = f"tdb-table@ph{ph.placeholder_format.idx}"
        remove_shape(ph)
        log.warn(n, "table placed in the geometry of a text-only BODY placeholder")
    table = gf.table
    table.first_row = bool(header)
    total_w = gf.width
    ratios = spec.get("col_widths")
    if not ratios:
        lens = [max(len(str(r[c])) if c < len(r) else 0 for r in data) for c in range(nc)]
        ratios = [min(max(L, 6), 40) for L in lens]
    s = float(sum(ratios))
    widths = [int(total_w * r / s) for r in ratios]
    widths[-1] = total_w - sum(widths[:-1])
    for c, w in enumerate(widths):
        table.columns[c].width = Emu(w)
    size = spec.get("font_size") or (14 if nr <= 5 else 12 if nr <= 8 else 11 if nr <= 11 else 10)
    row_h = int(gf.height / nr)
    hl = spec.get("highlight_row")
    for ri, row in enumerate(data):
        table.rows[ri].height = Emu(row_h)
        for ci in range(nc):
            val = str(row[ci]) if ci < len(row) else ""
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            _runs_from_markup(p, val)
            for r in p.runs:
                r.font.size = Pt(size)
                if hl is not None and ri == hl + (1 if header else 0):
                    r.font.bold = True
            if ci > 0 and (NUMERIC_RE.match(val) or val in ("-", "n/a")):
                p.alignment = PP_ALIGN.RIGHT
    return gf


# ---------------------------------------------------------------- slides

def columns_of(phs, slide_w):
    """Group content placeholders into columns by left edge (within 1/20 slide width)."""
    cols = []
    for ph in sorted(phs, key=lambda p: (rect(p)[0], rect(p)[1])):
        for col in cols:
            if abs(rect(col[0])[0] - rect(ph)[0]) < slide_w / 20:
                col.append(ph)
                break
        else:
            cols.append([ph])
    return [sorted(c, key=lambda p: rect(p)[1]) for c in cols]


def fill_body(slide, ph, spec, n, log):
    if "chart" in spec:
        return insert_chart(slide, ph, spec["chart"], n, log)
    if "table" in spec:
        return insert_table(slide, ph, spec["table"], n, log)
    items = spec.get("bullets", spec.get("body"))
    write_text(ph, items)
    return ph


def fill_slide(prs, slide, spec, n, log, edit=False):
    W, H = prs.slide_width, prs.slide_height
    used = set()
    info = {"slide": n, "layout": slide.slide_layout.name, "filled": [], "added_shapes": []}

    def mark(ph, what):
        used.add(ph.placeholder_format.idx)
        info["filled"].append(f"ph{ph.placeholder_format.idx}:{what}")

    # explicit placeholder overrides by idx
    for key, val in (spec.get("placeholders") or {}).items():
        try:
            ph = slide.placeholders[int(key)]
        except KeyError:
            log.err(n, f"placeholder idx {key} not on layout '{slide.slide_layout.name}'")
            continue
        mark(ph, "explicit")  # before filling: insert_chart/table invalidates ph
        if isinstance(val, dict) and ("chart" in val or "table" in val):
            fill_body(slide, ph, val, n, log)
        else:
            write_text(ph, val, preserve=edit)

    if "title" in spec:
        tph = title_placeholder(slide)
        if tph is None:
            log.err(n, f"spec has a title but layout '{slide.slide_layout.name}' has no title placeholder")
        elif tph.placeholder_format.idx not in used:
            write_text(tph, spec["title"], preserve=edit)
            mark(tph, "title")

    content = [p for p in content_placeholders(slide, H) if p.placeholder_format.idx not in used]
    subs = [p for p in content if ph_type(p) == PP_PLACEHOLDER.SUBTITLE]
    bodies = [p for p in content if ph_type(p) != PP_PLACEHOLDER.SUBTITLE]
    bodies.sort(key=lambda p: -(rect(p)[2] * rect(p)[3]))

    if "subtitle" in spec:
        target = (subs or bodies or [None])[0]
        if target is None:
            log.err(n, "subtitle given but no subtitle/body placeholder on this layout")
        else:
            write_text(target, spec["subtitle"], preserve=edit)
            mark(target, "subtitle")
            bodies = [b for b in bodies if b is not target]

    if "quote" in spec:
        if not bodies:
            log.err(n, "quote given but layout has no body placeholder")
        else:
            q = spec["quote"].strip().strip('"\u201c\u201d')
            write_text(bodies[0], "\u201c" + q + "\u201d", preserve=edit)
            mark(bodies[0], "quote")
            if "attribution" in spec:
                if len(bodies) > 1:
                    write_text(bodies[1], spec["attribution"], preserve=edit)
                    mark(bodies[1], "attribution")
                else:
                    log.err(n, "attribution given but layout has only one body placeholder")
            bodies = bodies[2:]

    if "columns" in spec:
        cols = columns_of(bodies, W)
        if len(cols) < len(spec["columns"]):
            log.err(n, f"{len(spec['columns'])} columns requested, layout '{slide.slide_layout.name}' has {len(cols)}; pick a two-column layout")
        for col_spec, col in zip(spec["columns"], cols):
            if len(col) >= 2 and "heading" in col_spec:
                write_text(col[0], col_spec["heading"], preserve=edit)
                mark(col[0], "column heading")
                mark(col[1], "column body")
                fill_body(slide, col[1], col_spec, n, log)
            else:
                cs = dict(col_spec)
                if "heading" in cs and "chart" not in cs and "table" not in cs:
                    cs["bullets"] = [f"**{cs['heading']}**"] + [[b] if isinstance(b, str) else b for b in _as_list(cs.get("bullets"))]
                mark(col[0], "column")
                fill_body(slide, col[0], cs, n, log)
        bodies = []

    if any(k in spec for k in ("chart", "table", "bullets", "body")):
        if not bodies:
            log.err(n, f"content given but layout '{slide.slide_layout.name}' has no free content placeholder")
        else:
            target = bodies[0]
            if "chart" in spec and ph_type(target) not in CHART_OK:
                better = [b for b in bodies if ph_type(b) in CHART_OK]
                target = better[0] if better else target
            if "table" in spec and ph_type(target) not in TABLE_OK:
                better = [b for b in bodies if ph_type(b) in TABLE_OK]
                target = better[0] if better else target
            if ("bullets" in spec or "body" in spec) and ph_type(target) in (PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.TABLE):
                log.err(n, "text given but the content placeholder is chart/table only; use a Title and Content layout")
            else:
                mark(target, "chart" if "chart" in spec else "table" if "table" in spec else "bullets")
                fill_body(slide, target, spec, n, log)

    if spec.get("source"):
        src = [p for p in slide.placeholders if is_source_ph(p, H) and p.placeholder_format.idx not in used]
        if src:
            write_text(src[0], spec["source"], preserve=edit)
            mark(src[0], "source")
        else:
            add_source_box(prs, slide, spec["source"])
            info["added_shapes"].append("tdb-source")
            log.warn(n, "layout has no source placeholder; added text box 'tdb-source' aligned to the content edge")

    for ft in spec.get("free_text") or []:
        tb = slide.shapes.add_textbox(Inches(ft["left_in"]), Inches(ft["top_in"]),
                                      Inches(ft["width_in"]), Inches(ft["height_in"]))
        tb.name = "tdb-free"
        tb.text_frame.word_wrap = True
        write_text(tb, ft["text"], preserve=False)
        if ft.get("size_pt"):
            for p in tb.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(ft["size_pt"])
        info["added_shapes"].append("tdb-free")
        log.warn(n, "free-floating text box added because the spec asked for it (free_text)")

    if "notes" in spec:
        slide.notes_slide.notes_text_frame.text = spec["notes"]

    if not spec.get("keep_empty_placeholders", False):
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx in used or ph_type(ph) in TITLE_TYPES and "title" in spec:
                continue
            if ph.has_text_frame and not ph.text_frame.text.strip():
                info.setdefault("pruned", []).append(f"ph{ph.placeholder_format.idx}:{ph.name}")
                remove_shape(ph)
    log.slides.append(info)


def _as_list(x):
    return [] if x is None else ([x] if isinstance(x, str) else list(x))


def add_source_box(prs, slide, text):
    H = prs.slide_height
    tph = title_placeholder(slide)
    left = rect(tph)[0] if tph is not None and rect(tph) else Inches(0.5)
    width = (rect(tph)[2] if tph is not None and rect(tph) else prs.slide_width - 2 * left)
    bottoms = [rect(s)[1] + rect(s)[3] for s in slide.shapes if rect(s) and s is not tph]
    # bottom-left, below everything else on the slide, never off the slide
    top = min(max(max(bottoms or [0]) + Inches(0.05), H - Inches(0.55)), H - Inches(0.40))
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Inches(0.32))
    tb.name = "tdb-source"
    tb.text_frame.word_wrap = True
    write_text(tb, text, preserve=False)
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.size = Pt(10)
        r.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
        r.font.color.brightness = 0.35


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if prs.part.related_part(sldId.rId) is slide.part:
            prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
            return


def recommended_layouts(prs):
    best = {}
    for layout in (l for m in prs.slide_masters for l in m.slide_layouts):
        _, scores = classify(layout, prs.slide_width, prs.slide_height)
        for k, v in scores.items():
            if v >= 5 and v > best.get(k, (0, None))[0]:
                best[k] = (v, layout)
    return {k: v[1] for k, v in best.items()}


TYPE_ALIASES = {"chart": "one_message_chart", "content": "one_message_chart", "bullets": "agenda",
                "cover": "title", "divider": "section", "two-column": "two_column"}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("storyboard")
    ap.add_argument("--template", help="company .pptx/.potx (overrides storyboard 'template')")
    ap.add_argument("--edit", help="existing deck to edit in place (keeps masters, layouts, slides)")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()

    sb = json.loads(Path(a.storyboard).read_text())
    base = a.edit or a.template or sb.get("template")
    if not base:
        sys.exit("No template: pass --template or set 'template' in the storyboard")
    prs = open_presentation(base)
    edit = bool(a.edit)
    original = list(prs.slides)
    if not edit and not sb.get("keep_template_slides", False):
        for s in original:
            delete_slide(prs, s)
        original = []
    rec = recommended_layouts(prs)
    log = Log()

    for i, spec in enumerate(sb.get("slides", []), 1):
        if "edit" in spec:
            k = int(spec["edit"])
            if not 1 <= k <= len(original):
                log.err(i, f"edit target slide {k} does not exist (deck has {len(original)})")
                continue
            fill_slide(prs, original[k - 1], spec, k, log, edit=True)
            continue
        try:
            if "layout" in spec:
                layout = find_layout(prs, spec["layout"])
            else:
                t = TYPE_ALIASES.get(spec.get("type"), spec.get("type"))
                layout = rec.get(t)
                if layout is None:
                    raise KeyError(f"no layout given and none suits type '{spec.get('type')}'")
        except KeyError as e:
            log.err(i, str(e).strip('"'))
            continue
        slide = prs.slides.add_slide(layout)
        fill_slide(prs, slide, spec, len(prs.slides), log)

    for k in sorted({int(x) for x in sb.get("delete", [])}, reverse=True):
        if 1 <= k <= len(original):
            delete_slide(prs, original[k - 1])
    if sb.get("order"):
        lst = prs.slides._sldIdLst
        ids = list(lst)
        for el in ids:
            lst.remove(el)
        for k in sb["order"]:
            lst.append(ids[k - 1])

    prs.save(a.out)
    log_path = Path(a.out).with_suffix(".build_log.json")
    log_path.write_text(json.dumps({"slides": log.slides, "errors": log.errors, "warnings": log.warnings}, indent=2))
    for s in log.slides:
        extra = f" added={s['added_shapes']}" if s["added_shapes"] else ""
        pr = f" pruned={s['pruned']}" if s.get("pruned") else ""
        print(f"  slide {s['slide']}: [{s['layout']}] {', '.join(s['filled'])}{extra}{pr}")
    for w in log.warnings:
        print("WARN  " + w)
    for e in log.errors:
        print("ERROR " + e)
    print(f"Wrote {a.out} ({len(prs.slides)} slides); log {log_path}")
    sys.exit(1 if log.errors else 0)


if __name__ == "__main__":
    main()
