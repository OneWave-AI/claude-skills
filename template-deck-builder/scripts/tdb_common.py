"""Shared helpers for template-deck-builder scripts (python-pptx >= 0.6.21).

Kept in one module so inspect/build/qa agree on how a template is opened,
how theme values are read, and how placeholders are classified.
"""
import io
import re
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

EMU_PER_INCH = 914400
EMU_PER_PT = 12700

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}

POTX_CT = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PPTX_CT = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"

# Placeholder types that hold the slide's main content (text, chart, table, picture).
CONTENT_TYPES = {
    PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.CHART,
    PP_PLACEHOLDER.TABLE, PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.BITMAP,
    PP_PLACEHOLDER.ORG_CHART, PP_PLACEHOLDER.MEDIA_CLIP,
}
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
CHROME_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
SOURCE_NAME_RE = re.compile(r"source|footnote|note", re.I)


def open_presentation(path):
    """Open .pptx or .potx. python-pptx rejects the .potx main content type,
    so rewrite [Content_Types].xml in memory; nothing else differs."""
    if str(path).lower().endswith((".potx", ".potm")):
        buf_in = open(path, "rb").read()
        zin = zipfile.ZipFile(io.BytesIO(buf_in))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(POTX_CT, PPTX_CT)
                zout.writestr(item, data)
        out.seek(0)
        return Presentation(out)
    return Presentation(path)


def emu_to_in(v):
    return None if v is None else round(v / EMU_PER_INCH, 3)


def read_theme(master):
    """Return {'colors': {dk1: 'RRGGBB', ...}, 'fonts': {'major': .., 'minor': ..}}."""
    theme_part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    colors = {}
    scheme = root.find(".//a:clrScheme", NS)
    if scheme is not None:
        for child in scheme:
            tag = etree.QName(child).localname
            srgb = child.find("a:srgbClr", NS)
            sysc = child.find("a:sysClr", NS)
            if srgb is not None:
                colors[tag] = srgb.get("val").upper()
            elif sysc is not None:
                colors[tag] = (sysc.get("lastClr") or "000000").upper()
    fonts = {}
    for kind in ("major", "minor"):
        el = root.find(f".//a:fontScheme/a:{kind}Font/a:latin", NS)
        fonts[kind] = el.get("typeface") if el is not None else None
    return {"name": root.get("name"), "colors": colors, "fonts": fonts}


def ph_type(ph):
    try:
        return ph.placeholder_format.type
    except Exception:
        return None


def rect(shape):
    """(left, top, width, height) in EMU; placeholders inherit from layout."""
    try:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if None in (l, t, w, h):
        return None
    return (int(l), int(t), int(w), int(h))


def is_source_ph(ph, slide_h):
    """A small body placeholder near the bottom, or one named Source/Footnote."""
    if ph_type(ph) not in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
        return False
    if SOURCE_NAME_RE.search(ph.name or ""):
        return True
    r = rect(ph)
    return bool(r and r[3] < slide_h * 0.08 and r[1] > slide_h * 0.8)


def content_placeholders(container, slide_h):
    """Content placeholders excluding source-line ones, sorted left-to-right, top-to-bottom."""
    out = []
    for ph in container.placeholders:
        t = ph_type(ph)
        if t in CONTENT_TYPES or t == PP_PLACEHOLDER.SUBTITLE:
            if t != PP_PLACEHOLDER.SUBTITLE and is_source_ph(ph, slide_h):
                continue
            out.append(ph)
    return sorted(out, key=lambda p: ((rect(p) or (0, 0))[0] // (EMU_PER_INCH // 4), (rect(p) or (0, 0, 0, 0))[1]))


def title_placeholder(container):
    for ph in container.placeholders:
        if ph_type(ph) in TITLE_TYPES:
            return ph
    return None


def find_layout(prs, name=None, index=None):
    layouts = [l for m in prs.slide_masters for l in m.slide_layouts]
    if index is not None:
        return layouts[index]
    for l in layouts:
        if l.name == name:
            return l
    for l in layouts:
        if l.name.strip().lower() == str(name).strip().lower():
            return l
    raise KeyError(f"Layout '{name}' not in template. Available: {[l.name for l in layouts]}")
