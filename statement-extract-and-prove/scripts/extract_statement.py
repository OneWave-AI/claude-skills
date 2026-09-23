#!/usr/bin/env python3
"""Deterministic transaction extraction from bank, card and brokerage statement PDFs.

Reads word positions with pdfplumber, finds the column header on each page, builds
x-bands per column, and assembles rows line by line. Every output row carries its page
number and y-position so any number can be traced back to the PDF. Nothing is guessed:
anything parsed with a fallback is flagged, and every line inside the table region that
did not become a row is written to the meta file as an unassigned line.

Usage:
  python extract_statement.py statement.pdf --out out/stmt
      [--account-type bank|card] [--decimal auto|dot|comma] [--date-order auto|mdy|dmy]
      [--year 2026] [--bands "date=40-95,description=95-330,amount=330-430,balance=430-560"]
      [--band-pages 2,3] [--allow-integers] [--xlsx]

Outputs:
  <out>.csv        one row per transaction (feeds prove.py and bookkeeping-close reconcile.py)
  <out>.meta.json  statement summary, markers, bands used per page, warnings, unassigned lines
  <out>.xlsx       optional, same rows plus a Summary sheet

Exit codes: 0 ok, 2 bad input/ambiguous locale, 3 no text layer (OCR first, see references/ocr.md).
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    sys.exit("pdfplumber is required: pip install pdfplumber")

# ---------------------------------------------------------------- vocab

ROLE_PATTERNS = [
    ("balance", r"(running )?balance( \(\w+\))?|saldo|solde|kontostand|balance owed"),
    ("debit", r"debits?|withdrawals?( and debits)?|money out|paid out|out|soll|belastung|lastschrift|"
              r"d[ée]bit|debit amount|charges|withdrawal amount"),
    ("credit", r"credits?|deposits?( and credits)?|money in|paid in|in|haben|gutschrift|cr[ée]dit|"
               r"credit amount|deposit amount"),
    ("amount", r"amount( \(\w+\))?|betrag( \(\w+\))?|montant|importe|umsatz|net amount|transaction amount"),
    ("date", r"(trans(action)?\.? |posting |post |booking |value |settle(ment)? |trade )?date|datum|"
             r"fecha|posted|buchungstag|valuta|wert(stellung)?|date op[ée]ration|date valeur"),
    ("description", r"description|details|transaction( details)?|particulars|memo|payee|narrative|"
                    r"buchungstext|verwendungszweck|libell[ée]|concepto|vorgang|activity|merchant"),
    ("reference", r"ref(erence)?\.?( no\.?)?|check( no\.?| #)?|cheque( no\.?)?|doc(ument)? no\.?|"
                  r"receipt|confirmation"),
]
ROLE_RES = [(r, re.compile(rf"^(?:{p})$", re.I)) for r, p in ROLE_PATTERNS]
NUMERIC_ROLES = ("debit", "credit", "amount", "balance")
TEXT_ROLES = ("description", "reference")

OPENING_RE = re.compile(r"([üu]bertrag (von|aus)|balance brought forward|brought forward|balance forward|opening balance|"
                        r"beginning balance|previous balance|starting balance|saldo vortrag|alter saldo|"
                        r"anfangssaldo|solde (pr[ée]c[ée]dent|initial|report[ée])|\bb/f\b)", re.I)
CARRIED_RE = re.compile(r"(balance carried forward|carried forward|\bc/f\b|saldo [üu]bertrag|"
                        r"[üu]bertrag( (auf|nach))?|solde [àa] reporter)", re.I)
CLOSING_RE = re.compile(r"(closing balance|ending balance|new balance|final balance|neuer saldo|"
                        r"endsaldo|solde (final|nouveau))", re.I)
TOTAL_RE = re.compile(r"^(total|totals|sub-?total|page total|summe|gesamt|totaux?)\b", re.I)
FOOTER_RE = re.compile(r"(page \d+( of \d+)?|seite \d+( von \d+)?|continued (on|from)|"
                       r"member fdic|\(continued\)|fortsetzung|suite au verso)", re.I)

PERIOD_RE = re.compile(r"(?:statement period|period|for the period|from|zeitraum|p[ée]riode|du)\s*:?\s*"
                       r"(?P<a>[\w./,\- ]{5,20}?)\s*(?:-|–|to|through|thru|bis|au)\s*"
                       r"(?P<b>[\w./, ]{5,20}\d{2,4})", re.I)
SUMMARY_PATTERNS = {
    "opening_balance": r"(beginning|opening|previous|starting) balance|anfangssaldo|alter saldo|"
                       r"solde (pr[ée]c[ée]dent|initial)",
    "closing_balance": r"(ending|closing|new|final) balance|neuer saldo|endsaldo|solde (final|nouveau)",
    "total_credits": r"total (deposits|credits|paid in|money in)|deposits (and|&) (other )?(additions|credits)|"
                     r"payments (and|&) (other )?credits|summe haben|summe gutschriften|total cr[ée]dits?",
    "total_debits": r"total (withdrawals|debits|paid out|money out)|withdrawals (and|&) (other )?(subtractions|debits)|"
                    r"purchases (and|&) (other )?(charges|debits)|summe soll|summe belastungen|"
                    r"total d[ée]bits?",
    "transaction_count": r"(number of|total) transactions|anzahl( der)? buchungen|nombre d'op[ée]rations",
}
SUMMARY_RES = {k: re.compile(v, re.I) for k, v in SUMMARY_PATTERNS.items()}

MONTHS = {"jan": 1, "feb": 2, "fev": 2, "fév": 2, "mar": 3, "mär": 3, "mrz": 3, "apr": 4, "avr": 4,
          "may": 5, "mai": 5, "jun": 6, "juin": 6, "jul": 7, "juil": 7, "aug": 8, "aoû": 8, "aou": 8,
          "sep": 9, "oct": 10, "okt": 10, "nov": 11, "dec": 12, "dez": 12, "déc": 12}

CURRENCY_RE = re.compile(r"[$€£¥]|\b(USD|EUR|GBP|CHF|CAD|AUD)\b")
MONEY_RE = re.compile(
    r"^(?P<lp>\()?(?P<sign>[-+−])?(?P<num>\d{1,3}(?:[.,'   ]\d{3})+(?:[.,]\d{1,3})?|\d+(?:[.,]\d{1,3})?)"
    r"(?P<rp>\))?\s?(?P<suf>-|\+|CR|DR|Cr|Dr|cr|dr|H|S)?$")
# date_inherited is recorded but not low-confidence: many layouts print the date once per day.
LOW_CONFIDENCE_FLAGS = {"no_amount", "locale_conflict", "bad_grouping", "text_in_numeric_band",
                        "collision", "negative_in_directional_column", "both_debit_and_credit",
                        "missing_year", "unparsed_date", "amount_on_continuation", "integer_amount"}


@dataclass
class Word:
    text: str
    x0: float
    x1: float
    top: float
    bottom: float

    @property
    def xc(self):
        return (self.x0 + self.x1) / 2


@dataclass
class Row:
    page: int
    y: float
    date_raw: str = ""
    post_date_raw: str = ""
    description: list = field(default_factory=list)
    reference: list = field(default_factory=list)
    nums: dict = field(default_factory=dict)  # role -> raw text
    flags: set = field(default_factory=set)
    last_y: float = 0.0


# ---------------------------------------------------------------- geometry

def page_lines(page, tol=2.5):
    words = page.extract_words(x_tolerance=1.5, y_tolerance=2, keep_blank_chars=False, use_text_flow=False)
    ws = sorted((Word(w["text"], w["x0"], w["x1"], w["top"], w["bottom"]) for w in words),
                key=lambda w: (round(w.top, 1), w.x0))
    lines = []
    for w in ws:
        if lines and abs(w.top - lines[-1][0].top) <= tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [merge_tokens(sorted(l, key=lambda w: w.x0)) for l in lines]


def merge_tokens(line):
    """Glue split numbers: '1 234,56', '$ 12.00', '12.00 CR', '12.00 -'."""
    out = []
    for w in line:
        if out:
            p = out[-1]
            gap = w.x0 - p.x1
            cw = max((p.x1 - p.x0) / max(len(p.text), 1), 3.0)
            close = gap < cw * 1.2
            if close and (
                (re.fullmatch(r"[-+]?\(?\d{1,3}", p.text) and re.fullmatch(r"\d{3}([.,]\d{1,3})?\)?(-|CR|DR)?", w.text))
                or re.fullmatch(r"[$€£¥]|-[$€£¥]|\(|-", p.text) and MONEY_RE.match(clean_currency(w.text))
                or (MONEY_RE.match(clean_currency(p.text)) and re.fullmatch(r"CR|DR|Cr|Dr|-", w.text))
            ):
                sep = "" if re.fullmatch(r"[$€£¥(-]|-[$€£¥]", p.text) or w.text in ("-",) else " "
                if sep == " " and re.fullmatch(r"\d{3}.*", w.text):
                    sep = " "  # thousands group separated by a space
                out[-1] = Word(p.text + sep + w.text, p.x0, w.x1, min(p.top, w.top), max(p.bottom, w.bottom))
                continue
        out.append(w)
    return out


def clean_currency(s):
    return CURRENCY_RE.sub("", s).strip()


def is_money(text, allow_integers=False):
    m = MONEY_RE.match(clean_currency(text))
    if not m:
        return False
    num = m.group("num")
    has_dec = re.search(r"[.,]\d{1,2}$", num) is not None
    return has_dec or allow_integers


def phrases(line):
    """Group words into phrases split at gaps wider than ~a space."""
    out = []
    for w in line:
        if out:
            p = out[-1]
            cw = (p[-1].x1 - p[-1].x0) / max(len(p[-1].text), 1)
            if w.x0 - p[-1].x1 < max(cw * 1.6, 4.0):
                p.append(w)
                continue
        out.append([w])
    return [Word(" ".join(x.text for x in p), p[0].x0, p[-1].x1, p[0].top, p[-1].bottom) for p in out]


def detect_header(line):
    cols = []
    def role_of(text):
        t = re.sub(r"[:*]+$", "", text.strip()).strip()
        return next((r for r, rx in ROLE_RES if rx.match(t)), None)

    for ph in phrases(line):
        role = role_of(ph.text)
        if role:
            cols.append((role, ph))
        elif " " in ph.text:  # tightly set headers merge into one phrase: try word by word
            ws = [w for w in line if w.x0 >= ph.x0 - 0.1 and w.x1 <= ph.x1 + 0.1]
            for w in ws:
                r = role_of(w.text)
                if r:
                    cols.append((r, w))
    roles = [r for r, _ in cols]
    if "date" in roles and any(r in NUMERIC_ROLES for r in roles) and len(set(roles)) >= 3:
        # second date column becomes post_date
        seen, fixed = set(), []
        for r, ph in cols:
            if r == "date" and "date" in seen:
                r = "post_date"
            seen.add(r)
            fixed.append((r, ph))
        return fixed
    return None


def bands_from_header(cols):
    b = {r: (ph.x0, ph.x1) for r, ph in cols}
    return b


def parse_bands_arg(s):
    b = {}
    for part in s.split(","):
        k, v = part.split("=")
        a, z = v.split("-")
        b[k.strip()] = (float(a), float(z))
    return b


def assign(line, bands, explicit, allow_integers):
    """Return dict role -> list[Word] for one line."""
    got = {}
    num_roles = [r for r in NUMERIC_ROLES if r in bands]
    date_roles = [r for r in ("date", "post_date") if r in bands]
    text_roles = [r for r in TEXT_ROLES if r in bands]
    num_left = min((bands[r][0] for r in num_roles), default=1e9)
    if explicit:
        for w in line:
            role = next((r for r, (a, z) in bands.items() if a <= w.xc <= z), None)
            if role in NUMERIC_ROLES and not is_money(w.text, allow_integers):
                role = "_stray"
            got.setdefault(role or "_outside", []).append(w)
        return got
    first_text_x0 = min((bands[r][0] for r in text_roles), default=None)
    for w in line:
        if is_money(w.text, allow_integers) and w.x1 >= num_left - 1:
            def score(r):
                a, z = bands[r]
                d = 0 if (w.x0 <= z and w.x1 >= a) else min(abs(w.x0 - z), abs(w.x1 - a))
                return (d, abs(w.x1 - z))
            role = min(num_roles, key=score)
        elif w.x1 >= num_left - 1 and w.x0 > num_left - 1:
            role = "_stray"
        elif date_roles and first_text_x0 is not None and w.x0 < first_text_x0 - 1:
            # date zone: nearest date column by left edge
            role = min(date_roles, key=lambda r: abs(w.x0 - bands[r][0]))
        else:
            if text_roles:
                role = min(text_roles, key=lambda r: 0 if bands[r][0] - 2 <= w.x0 else abs(w.x0 - bands[r][0]))
                # a word left of the reference band but right of description start belongs to the nearest-left band
                left = [r for r in text_roles if bands[r][0] - 2 <= w.x0]
                if left:
                    role = max(left, key=lambda r: bands[r][0])
            else:
                role = "description"
        got.setdefault(role, []).append(w)
    return got


# ---------------------------------------------------------------- numbers

def detect_decimal(tokens):
    dot = comma = 0
    for t in tokens:
        m = MONEY_RE.match(clean_currency(t))
        if not m:
            continue
        num = m.group("num")
        if re.search(r"\.\d{2}$", num) and not re.search(r",\d{2}$", num):
            dot += 1
        elif re.search(r",\d{2}$", num):
            comma += 1
    return dot, comma


def parse_amount(raw, decimal, account_type):
    """Return (Decimal value with printed sign applied, set of flags)."""
    flags = set()
    s = clean_currency(raw).replace("−", "-")
    m = MONEY_RE.match(s)
    if not m:
        return None, {"unparsed_amount"}
    num = m.group("num")
    dec_sep, thou_sep = (".", ",") if decimal == "dot" else (",", ".")
    tail = re.search(r"[.,](\d{1,3})$", num)
    if tail and num[tail.start()] != dec_sep and len(tail.group(1)) != 3:
        flags.add("locale_conflict")  # e.g. '12,34' in a dot-decimal document
    if tail and num[tail.start()] == dec_sep:
        int_part, frac = num[:tail.start()], tail.group(1)
    else:
        int_part, frac = num, ""
        if not re.search(r"[.,]\d{2}$", num):
            flags.add("integer_amount")
    groups = re.split(r"[.,'   ]", int_part)
    if len(groups) > 1:
        seps = set(re.findall(r"[.,'   ]", int_part))
        if dec_sep in seps or any(len(g) != 3 for g in groups[1:]) or len(groups[0]) > 3:
            flags.add("bad_grouping")
    digits = "".join(groups)
    try:
        val = Decimal(digits + ("." + frac if frac else ""))
    except InvalidOperation:
        return None, {"unparsed_amount"}
    neg = bool(m.group("lp") and m.group("rp")) or m.group("sign") in ("-", "−") or m.group("suf") == "-"
    suf = (m.group("suf") or "").upper()
    if suf in ("CR", "H"):  # H = Haben (German credit)
        neg = account_type == "card"
    elif suf in ("DR", "S"):  # S = Soll (German debit)
        neg = account_type != "card"
    return (-val if neg else val), flags


# ---------------------------------------------------------------- dates

def parse_date(s, order, year_hint=None):
    """Return (date|None, had_year)."""
    s = s.strip().rstrip(".").strip()
    if not s:
        return None, False
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
        return safe_date(y, mo, d), True
    m = re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})(?:[/.\-](\d{2,4}))?", s)
    if m:
        a, b, c = m.group(1), m.group(2), m.group(3)
        a, b = int(a), int(b)
        mo, d = (a, b) if order == "mdy" else (b, a)
        y = norm_year(c) if c else year_hint
        return safe_date(y, mo, d) if y else None, bool(c)
    m = re.fullmatch(r"(\d{1,2})\.?[\s\-]*([A-Za-zÀ-ÿ]{3,9})\.?,?[\s\-]*(\d{2,4})?", s)
    if m and month_num(m.group(2)):
        y = norm_year(m.group(3)) if m.group(3) else year_hint
        return safe_date(y, month_num(m.group(2)), int(m.group(1))) if y else None, bool(m.group(3))
    m = re.fullmatch(r"([A-Za-zÀ-ÿ]{3,9})\.?\s+(\d{1,2}),?\s*(\d{2,4})?", s)
    if m and month_num(m.group(1)):
        y = norm_year(m.group(3)) if m.group(3) else year_hint
        return safe_date(y, month_num(m.group(1)), int(m.group(2))) if y else None, bool(m.group(3))
    return None, False


def month_num(t):
    t = t.lower()
    for k in (t[:4], t[:3]):
        if k in MONTHS:
            return MONTHS[k]
    return None


def norm_year(c):
    y = int(c)
    return y + 2000 if y < 100 else y


def safe_date(y, mo, d):
    try:
        return date(y, mo, d)
    except (ValueError, TypeError):
        return None


def looks_like_date(s):
    s = s.strip()
    return bool(re.fullmatch(r"\d{1,4}[/.\-]\d{1,2}([/.\-]\d{2,4})?\.?", s)
                or re.fullmatch(r"\d{1,2}\.?\s*[A-Za-zÀ-ÿ]{3,9}\.?,?\s*(\d{2,4})?", s) and month_num(re.sub(r"[\d.,\s]", "", s))
                or re.fullmatch(r"[A-Za-zÀ-ÿ]{3,9}\.?\s+\d{1,2},?\s*(\d{2,4})?", s) and month_num(s.split()[0]))


def detect_date_order(samples):
    first_big = any(int(a) > 12 for a, b in samples)
    second_big = any(int(b) > 12 for a, b in samples)
    if first_big and second_big:
        return None
    if first_big:
        return "dmy"
    if second_big:
        return "mdy"
    return "ambiguous"


def infer_year(d_mo_day, period):
    """Pick the year that puts month/day inside the statement period (+/- 31 days)."""
    mo, dd = d_mo_day
    start, end = period
    best = None
    for y in {start.year - 1, start.year, end.year, end.year + 1}:
        c = safe_date(y, mo, dd)
        if not c:
            continue
        dist = 0 if start <= c <= end else min(abs((c - start).days), abs((c - end).days))
        if best is None or dist < best[0]:
            best = (dist, c)
    return best[1] if best and best[0] <= 31 else None


# ---------------------------------------------------------------- main extraction

def extract(pdf_path, a):
    pdf = pdfplumber.open(pdf_path)
    meta = {"file": str(pdf_path), "pages": len(pdf.pages), "account_type": a.account_type,
            "text_layer": [], "bands": {}, "markers": [], "unassigned_lines": [], "warnings": [],
            "summary": {}, "totals_rows": []}
    all_lines = []
    for i, page in enumerate(pdf.pages, 1):
        nchars = len(page.chars)
        meta["text_layer"].append({"page": i, "chars": nchars, "images": len(page.images)})
        all_lines.append(page_lines(page))
    total_chars = sum(p["chars"] for p in meta["text_layer"])
    empty = [p["page"] for p in meta["text_layer"] if p["chars"] < 20]
    if total_chars < 20 * len(pdf.pages) or len(empty) == len(pdf.pages):
        print(f"NO TEXT LAYER: {total_chars} characters across {len(pdf.pages)} pages. "
              f"This PDF is scanned or image-only. OCR it first (see references/ocr.md), e.g.\n"
              f"  ocrmypdf --deskew --rotate-pages -l eng {pdf_path} ocr.pdf", file=sys.stderr)
        meta["no_text_layer"] = True
        return None, meta, 3
    if empty:
        meta["warnings"].append(f"pages with no text layer (OCR these or read them manually): {empty}")

    explicit_bands = parse_bands_arg(a.bands) if a.bands else None
    band_pages = {int(x) for x in a.band_pages.split(",")} if a.band_pages else None

    # pass 1: headers + token roles per page
    page_struct = []
    bands = None
    for pi, lines in enumerate(all_lines, 1):
        hdr_idx, page_bands, source = None, bands, "carried from previous page"
        for li, line in enumerate(lines):
            h = detect_header(line)
            if h:
                hdr_idx, page_bands, source = li, bands_from_header(h), "header"
                break
        use_explicit = explicit_bands is not None and (band_pages is None or pi in band_pages)
        if use_explicit:
            page_bands, source = explicit_bands, "explicit --bands"
        if page_bands is None:
            meta["warnings"].append(f"page {pi}: no column header found and no bands to carry; page skipped")
            page_struct.append(None)
            continue
        bands = page_bands if source != "explicit --bands" else bands
        meta["bands"][pi] = {"source": source, "columns": {k: [round(v[0], 1), round(v[1], 1)]
                                                            for k, v in page_bands.items()}}
        assigned = [assign(l, page_bands, use_explicit, a.allow_integers) for l in lines]
        page_struct.append({"lines": lines, "hdr": hdr_idx, "assigned": assigned})

    # locale detection from tokens that landed in numeric columns only
    num_tokens = [w.text for ps in page_struct if ps for asg in ps["assigned"]
                  for r in NUMERIC_ROLES for w in asg.get(r, [])]
    dot, comma = detect_decimal(num_tokens)
    if not (dot or comma):  # no table found yet: fall back to every token on every page
        dot, comma = detect_decimal([w.text for lines in all_lines for l in lines for w in l])
    decimal = a.decimal
    if decimal == "auto":
        if dot and comma and min(dot, comma) / (dot + comma) > 0.1:
            print(f"AMBIGUOUS DECIMAL SEPARATOR: {dot} tokens end '.dd', {comma} end ',dd'. "
                  f"Re-run with --decimal dot or --decimal comma.", file=sys.stderr)
            return None, meta, 2
        decimal = "comma" if comma > dot else "dot"
    meta["decimal"] = {"used": decimal, "requested": a.decimal, "evidence": {"dot": dot, "comma": comma}}

    # date order detection from date-zone tokens
    samples = []
    for ps in page_struct:
        if not ps:
            continue
        for asg in ps["assigned"]:
            t = " ".join(w.text for w in asg.get("date", []))
            m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})", t)
            if m:
                samples.append(m.groups())
    order = a.date_order
    if order == "auto":
        det = detect_date_order(samples)
        if det is None:
            print("DATE ORDER CONFLICT: some dates only parse as D/M, others only as M/D. "
                  "Check the date band; re-run with --date-order.", file=sys.stderr)
            return None, meta, 2
        if det == "ambiguous":
            order = "dmy" if decimal == "comma" else "mdy"
            if samples:
                meta["warnings"].append(f"date order not provable from data (no day > 12); assumed {order}. "
                                        f"Confirm against the statement period or pass --date-order.")
        else:
            order = det
    meta["date_order"] = order

    # summary + period from every line (non-transaction lines are the source of truth)
    period = None
    for pi, lines in enumerate(all_lines, 1):
        for line in lines:
            text = " ".join(w.text for w in line)
            if period is None:
                m = PERIOD_RE.search(text)
                if m:
                    d1, _ = parse_date(clean_date_text(m.group("a")), order)
                    d2, _ = parse_date(clean_date_text(m.group("b")), order)
                    if d1 and d2 and d1 <= d2:
                        period = (d1, d2)
            for key, rx in SUMMARY_RES.items():
                if key in meta["summary"] or not rx.search(text):
                    continue
                if key == "transaction_count":
                    m = re.search(r"(\d+)\s*$", text)
                    if m:
                        meta["summary"][key] = {"value": int(m.group(1)), "page": pi, "text": text}
                    continue
                money = [w for w in line if is_money(w.text, a.allow_integers)]
                if money:
                    v, fl = parse_amount(money[-1].text, decimal, a.account_type)
                    entry = {"value": str(v), "page": pi, "text": text}
                    cm = re.search(r"\((\d+)\)", text) or re.search(r"\b(\d+)\s+(items?|transactions?)\b", text, re.I)
                    if cm:
                        entry["count"] = int(cm.group(1))
                    meta["summary"][key] = entry
    if a.year and not period:
        period = (date(a.year, 1, 1), date(a.year, 12, 31))
    meta["period"] = [period[0].isoformat(), period[1].isoformat()] if period else None

    # pass 2: rows
    rows = []
    for pi, ps in enumerate(page_struct, 1):
        if not ps:
            continue
        lines, hdr, assigned = ps["lines"], ps["hdr"], ps["assigned"]
        pitch = median_pitch(lines)
        cur, ended = None, False
        start = (hdr + 1) if hdr is not None else 0
        for li in range(start, len(lines)):
            line, asg = lines[li], assigned[li]
            text = " ".join(w.text for w in line)
            y = round(line[0].top, 1)
            if hdr is not None and li > hdr and detect_header(line):
                ended, cur = False, None
                continue
            if FOOTER_RE.search(text):
                continue
            nums = {r: asg[r][0].text for r in NUMERIC_ROLES if asg.get(r)}
            date_txt = " ".join(w.text for w in asg.get("date", []))
            desc_txt = " ".join(w.text for w in asg.get("description", []))
            kind = marker_kind(text)
            if kind and not looks_like_date(date_txt):
                val_raw = nums.get("balance") or (list(nums.values())[-1] if nums else None)
                if kind == "total":
                    meta["totals_rows"].append({"page": pi, "y": y, "text": text, "values": {
                        r: str(parse_amount(t, decimal, a.account_type)[0]) for r, t in nums.items()}})
                    ended = True
                elif val_raw:
                    v, _ = parse_amount(val_raw, decimal, a.account_type)
                    if kind == "carried_forward" and not any(r.page == pi for r in rows):
                        kind = "brought_forward"  # a bare 'Ubertrag'/'forward' line above the first row
                    meta["markers"].append({"page": pi, "y": y, "kind": kind, "value": str(v), "text": text})
                    if kind == "closing":
                        ended = True
                cur = None
                continue
            if ended:
                meta["unassigned_lines"].append({"page": pi, "y": y, "text": text, "why": "after table end"})
                continue
            has_date = looks_like_date(date_txt)
            if has_date:
                cur = Row(page=pi, y=y, date_raw=date_txt, last_y=y)
                rows.append(cur)
                fill(cur, asg, nums)
                continue
            if nums:
                if cur is not None and not any(r in cur.nums for r in ("debit", "credit", "amount")) \
                        and line[0].top - cur.last_y <= pitch * 2.2:
                    cur.flags.add("amount_on_continuation")
                    fill(cur, asg, nums)
                    cur.last_y = line[0].top
                    continue
                prev_date = rows[-1].date_raw if rows else ""
                cur = Row(page=pi, y=y, date_raw=prev_date, last_y=y, flags={"date_inherited"})
                rows.append(cur)
                fill(cur, asg, nums)
                continue
            # text only
            if cur is not None and cur.page == pi and line[0].top - cur.last_y <= pitch * 2.2 and desc_txt:
                cur.description.append(desc_txt)
                cur.reference += [w.text for w in asg.get("reference", [])]
                for k in ("_stray", "_outside"):
                    if asg.get(k):
                        cur.flags.add("text_in_numeric_band")
                        cur.description += [w.text for w in asg[k]]
                cur.last_y = line[0].top
                continue
            if date_txt and not has_date:
                meta["unassigned_lines"].append({"page": pi, "y": y, "text": text, "why": "date band text not a date"})
                continue
            meta["unassigned_lines"].append({"page": pi, "y": y, "text": text, "why": "text line not adjacent to a row"})

    out_rows = finalize(rows, decimal, order, period, a, meta)
    return out_rows, meta, 0


def clean_date_text(s):
    return re.sub(r"^(on|the|vom|du|le)\s+", "", s.strip(), flags=re.I).strip(" ,.")


def marker_kind(text):
    if CARRIED_RE.search(text):
        return "carried_forward"
    if OPENING_RE.search(text):
        return "brought_forward"
    if CLOSING_RE.search(text):
        return "closing"
    if TOTAL_RE.search(text):
        return "total"
    return None


def median_pitch(lines):
    ys = sorted(l[0].top for l in lines)
    gaps = sorted(b - a for a, b in zip(ys, ys[1:]) if b - a > 1)
    return gaps[len(gaps) // 2] if gaps else 12.0


def fill(row, asg, nums):
    row.description += [" ".join(w.text for w in asg.get("description", []))] if asg.get("description") else []
    row.reference += [w.text for w in asg.get("reference", [])]
    if asg.get("post_date"):
        row.post_date_raw = " ".join(w.text for w in asg["post_date"])
    for r, t in nums.items():
        if r in row.nums:
            row.flags.add("collision")
        else:
            row.nums[r] = t
    for r in NUMERIC_ROLES:
        if len(asg.get(r, [])) > 1:
            row.flags.add("collision")
    for k in ("_stray", "_outside"):
        if asg.get(k):
            row.flags.add("text_in_numeric_band")
            row.description += [w.text for w in asg[k]]


def finalize(rows, decimal, order, period, a, meta):
    out = []
    for r in rows:
        flags = set(r.flags)
        d, had_year = parse_date(r.date_raw, order)
        if d is None and r.date_raw:
            m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})\.?$", r.date_raw.strip())
            md = None
            if m:
                a1, b1 = int(m.group(1)), int(m.group(2))
                md = (a1, b1) if order == "mdy" else (b1, a1)
            else:
                probe, _ = parse_date(r.date_raw, order, year_hint=2000)
                md = (probe.month, probe.day) if probe else None
            if md and period:
                d = infer_year(md, period)
                if d:
                    flags.add("year_inferred")
                else:
                    flags.add("missing_year")
            elif md:
                flags.add("missing_year")
            else:
                flags.add("unparsed_date")
        pd_, _ = parse_date(r.post_date_raw, order, year_hint=d.year if d else None) if r.post_date_raw else (None, False)
        vals = {}
        for role, raw in r.nums.items():
            v, fl = parse_amount(raw, decimal, a.account_type)
            flags |= fl
            vals[role] = v
        debit, credit, amount = vals.get("debit"), vals.get("credit"), vals.get("amount")
        if amount is None and (debit is not None or credit is not None):
            if debit is not None and credit is not None:
                flags.add("both_debit_and_credit")
            for v in (debit, credit):
                if v is not None and v < 0:
                    flags.add("negative_in_directional_column")
            inc = credit if a.account_type == "bank" else debit
            dec = debit if a.account_type == "bank" else credit
            amount = (inc or Decimal(0)) - (dec or Decimal(0))
        if amount is not None and a.invert_amount and "amount" in vals:
            amount = -amount
        if amount is None:
            flags.add("no_amount")
        desc = " ".join(x for x in r.description if x).strip()
        low = sorted(flags & LOW_CONFIDENCE_FLAGS)
        out.append({
            "page": r.page, "y": r.y,
            "date": d.isoformat() if d else "", "post_date": pd_.isoformat() if pd_ else "",
            "description": desc, "reference": " ".join(r.reference),
            "debit": fmt(debit), "credit": fmt(credit), "amount": fmt(amount),
            "balance": fmt(vals.get("balance")),
            "raw_amount": r.nums.get("amount") or " | ".join(filter(None, [r.nums.get("debit"), r.nums.get("credit")])),
            "raw_balance": r.nums.get("balance", ""), "date_raw": r.date_raw,
            "flags": ";".join(sorted(flags)), "low_confidence": "1" if low else "",
        })
    layout = "debit_credit" if any(("debit" in r.nums or "credit" in r.nums) for r in rows) else "signed_amount"
    meta["layout"] = layout
    meta["row_count"] = len(out)
    return out


def fmt(v):
    return "" if v is None else f"{v:.2f}" if v == v.quantize(Decimal("0.01")) else str(v)


FIELDS = ["page", "y", "date", "post_date", "description", "reference", "debit", "credit", "amount",
          "balance", "raw_amount", "raw_balance", "date_raw", "flags", "low_confidence"]


def write_outputs(rows, meta, out, xlsx):
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(f"{out}.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    Path(f"{out}.meta.json").write_text(json.dumps(meta, indent=2, default=str))
    if xlsx:
        try:
            from openpyxl import Workbook
        except ImportError:
            print("openpyxl not installed; skipped .xlsx", file=sys.stderr)
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"
        ws.append(FIELDS)
        for r in rows:
            ws.append([float(r[k]) if k in ("debit", "credit", "amount", "balance") and r[k] else r[k] for k in FIELDS])
        for col in ("G", "H", "I", "J"):
            for c in ws[col][1:]:
                c.number_format = "#,##0.00;[Red]-#,##0.00"
        ws.freeze_panes = "A2"
        s = wb.create_sheet("Summary")
        for k, v in meta["summary"].items():
            s.append([k, v.get("value"), v.get("page")])
        s.append(["period", " to ".join(meta["period"]) if meta.get("period") else ""])
        s.append(["decimal", meta["decimal"]["used"]])
        s.append(["rows", len(rows)])
        wb.save(f"{out}.xlsx")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf")
    ap.add_argument("--out", required=True, help="output path prefix (writes .csv, .meta.json, .xlsx)")
    ap.add_argument("--account-type", choices=["bank", "card"], default="bank",
                    help="bank: credits/deposits raise the balance; card: charges raise the balance owed")
    ap.add_argument("--decimal", choices=["auto", "dot", "comma"], default="auto")
    ap.add_argument("--date-order", choices=["auto", "mdy", "dmy"], default="auto")
    ap.add_argument("--year", type=int, help="year for dates printed without one, when no period is found")
    ap.add_argument("--bands", help="explicit column x-bands, e.g. date=40-95,description=95-330,amount=330-430,balance=430-560")
    ap.add_argument("--band-pages", help="apply --bands only to these pages, e.g. 2 or 2,3")
    ap.add_argument("--allow-integers", action="store_true", help="accept amounts without decimals (price lists, JPY)")
    ap.add_argument("--invert-amount", action="store_true", help="flip a signed Amount column printed from the other party's view")
    ap.add_argument("--xlsx", action="store_true")
    ap.add_argument("--dump-words", type=int, metavar="PAGE",
                    help="print every line on PAGE with word x-positions (to choose --bands), then exit")
    a = ap.parse_args()
    if a.dump_words:
        with pdfplumber.open(a.pdf) as pdf:
            page = pdf.pages[a.dump_words - 1]
            print(f"page {a.dump_words}: width {page.width:.0f}, height {page.height:.0f}")
            for line in page_lines(page):
                print(f"y={line[0].top:6.1f}  " + "  ".join(f"{w.text}[{w.x0:.0f}-{w.x1:.0f}]" for w in line))
        return
    rows, meta, code = extract(a.pdf, a)
    if code:
        Path(a.out).parent.mkdir(parents=True, exist_ok=True)
        Path(f"{a.out}.meta.json").write_text(json.dumps(meta, indent=2, default=str))
        sys.exit(code)
    write_outputs(rows, meta, a.out, a.xlsx)
    low = sum(1 for r in rows if r["low_confidence"])
    print(f"{len(rows)} rows from {meta['pages']} pages | layout={meta['layout']} decimal={meta['decimal']['used']} "
          f"date_order={meta['date_order']} period={meta['period']} | low-confidence rows={low} "
          f"| unassigned lines={len(meta['unassigned_lines'])}")
    for w in meta["warnings"]:
        print("WARNING:", w)
    print(f"wrote {a.out}.csv and {a.out}.meta.json" + (" and .xlsx" if a.xlsx else ""))


if __name__ == "__main__":
    main()
