#!/usr/bin/env python3
"""build_quote.py -- deterministic quote/estimate/bid builder for trades and field service.

Reads a YAML or JSON quote spec, prices every line with Decimal arithmetic, validates
the numbers, and writes:
  <number>-client.xlsx    client quote with live formulas (amounts, subtotals, tax, totals,
                          payment schedule)
  <number>-client.pdf     client quote PDF (only if reportlab is installed; skipped otherwise)
  <number>-internal.xlsx  internal cost sheet: every component, burden, contingency,
                          overhead, profit, margin vs target, effective hourly
  <number>-summary.json   machine-readable totals and warnings

Pricing math (per cost component):
  direct      = materials (qty x (1 + waste), rounded up to pack size) x unit cost
              | labor hours x wage x (1 + burden)
              | equipment / subs / fees qty x unit cost
  contingency = direct x contingency_pct
  overhead    = (direct + contingency) x overhead.pct_of_cost + labor hours x overhead.per_labor_hour
  loaded      = direct + contingency + overhead
  sell        = loaded / (1 - margin)      when pricing by target margin
              = loaded x (1 + markup)      when pricing by markup
  markup = margin / (1 - margin);  margin = markup / (1 + markup)

Usage:
  python build_quote.py spec.yaml --out ./out
  python build_quote.py spec.json --out ./out --no-pdf
  python build_quote.py spec.yaml --check            # validate and print totals only
  python build_quote.py spec.yaml --explain WALLS    # step-by-step hand-check of one line
  python build_quote.py --margin-to-markup 0.30      # conversion helper
  python build_quote.py --markup-to-margin 0.50

Requires openpyxl for xlsx. PyYAML only for .yaml specs. reportlab only for the PDF.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from decimal import ROUND_CEILING, ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

D = Decimal
CENT = D("0.01")
SHARE_Q = D("0.000001")
ZERO = D("0")
ONE = D("1")
COMP_TYPES = ["materials", "labor", "equipment", "subs", "fees"]


class SpecError(Exception):
    pass


# ---------------------------------------------------------------- numbers
def dec(v, where):
    if isinstance(v, bool) or v is None:
        raise SpecError(f"{where}: expected a number, got {v!r}")
    if isinstance(v, D):
        return v
    try:
        return D(str(v).replace(",", "").replace("$", "").strip())
    except InvalidOperation:
        raise SpecError(f"{where}: expected a number, got {v!r}")


def money(x):
    """Round half away from zero to cents -- matches Excel ROUND(x, 2)."""
    return x.quantize(CENT, rounding=ROUND_HALF_UP)


def pct(x, places=1):
    return f"{(x * 100).quantize(D(1).scaleb(-places), rounding=ROUND_HALF_UP)}%"


def fmt(x):
    return f"${money(x):,.2f}"


def margin_to_markup(m):
    if not (ZERO <= m < ONE):
        raise SpecError(f"margin must be >= 0 and < 1 (got {m}); 100% margin is impossible")
    return m / (ONE - m)


def markup_to_margin(mu):
    if mu < ZERO:
        raise SpecError(f"markup must be >= 0 (got {mu})")
    return mu / (ONE + mu)


# ---------------------------------------------------------------- spec loading
def load_spec(path):
    p = Path(path)
    text = p.read_text()
    if p.suffix.lower() in (".yaml", ".yml"):
        try:
            import yaml
        except ImportError:
            sys.exit("PyYAML is not installed. Run `pip install pyyaml` or convert the spec to JSON.")
        return yaml.safe_load(text)
    return json.loads(text, parse_float=D)


class Pricing:
    def __init__(self, spec):
        pr = spec.get("pricing") or {}
        has_m, has_mu = "target_margin" in pr, "target_markup" in pr
        if has_m == has_mu:
            raise SpecError("pricing: set exactly one of target_margin or target_markup (they are different numbers)")
        if has_m:
            self.method, self.value = "margin", dec(pr["target_margin"], "pricing.target_margin")
            self.margin, self.markup = self.value, margin_to_markup(self.value)
        else:
            self.method, self.value = "markup", dec(pr["target_markup"], "pricing.target_markup")
            self.markup, self.margin = self.value, markup_to_margin(self.value)
        self.by_type = {}
        for t, ov in (pr.get("by_component") or {}).items():
            if t not in COMP_TYPES:
                raise SpecError(f"pricing.by_component: unknown component type {t!r} (use {COMP_TYPES})")
            if ov.get("at_cost"):
                self.by_type[t] = ("at_cost", ZERO)
            elif "margin" in ov:
                m = dec(ov["margin"], f"pricing.by_component.{t}.margin")
                margin_to_markup(m)
                self.by_type[t] = ("margin", m)
            elif "markup" in ov:
                mu = dec(ov["markup"], f"pricing.by_component.{t}.markup")
                markup_to_margin(mu)
                self.by_type[t] = ("markup", mu)
            else:
                raise SpecError(f"pricing.by_component.{t}: give margin, markup, or at_cost: true")

    def rule(self, ctype):
        return self.by_type.get(ctype, (self.method, self.value))

    @staticmethod
    def sell(loaded, rule):
        kind, v = rule
        if kind == "margin":
            return loaded / (ONE - v)
        if kind == "markup":
            return loaded * (ONE + v)
        return loaded


class Config:
    def __init__(self, spec):
        self.spec = spec
        self.pricing = Pricing(spec)
        lab = spec.get("labor") or {}
        self.burden = dec(lab.get("burden_pct", 0), "labor.burden_pct")
        self.roles = {}
        for r, v in (lab.get("roles") or {}).items():
            if "wage" not in v:
                raise SpecError(f"labor.roles.{r}: missing wage (base hourly pay)")
            self.roles[r] = (dec(v["wage"], f"labor.roles.{r}.wage"),
                             dec(v.get("burden_pct", self.burden), f"labor.roles.{r}.burden_pct"))
        oh = spec.get("overhead") or {}
        self.oh_pct = dec(oh.get("pct_of_cost", 0), "overhead.pct_of_cost")
        self.oh_hr = dec(oh.get("per_labor_hour", 0), "overhead.per_labor_hour")
        self.cont = dec(spec.get("contingency_pct", 0), "contingency_pct")
        tax = spec.get("tax") or {"mode": "none"}
        self.tax_mode = tax.get("mode", "none")
        if self.tax_mode not in ("none", "charge_on_sell", "contractor_pays_on_cost"):
            raise SpecError("tax.mode must be none, charge_on_sell, or contractor_pays_on_cost")
        self.tax_rate = dec(tax.get("rate", 0), "tax.rate")
        self.tax_label = tax.get("label", "Sales tax")
        self.taxable = set(tax.get("taxable", ["materials"]))
        bad = self.taxable - set(COMP_TYPES)
        if bad:
            raise SpecError(f"tax.taxable: unknown types {sorted(bad)}")
        self.min_eff = dec(spec["min_effective_hourly"], "min_effective_hourly") if "min_effective_hourly" in spec else None
        self.client_view = (spec.get("client") or {}).get("view", "itemized")


# ---------------------------------------------------------------- line pricing
def expand_components(line, where):
    comps = []
    for t in COMP_TYPES:
        for i, c in enumerate(line.get(t) or []):
            comps.append((t, c, f"{where}.{t}[{i}]"))
    if "unit_cost" in line:  # shorthand: one component priced per display unit
        t = line.get("cost_type", "materials")
        c = {"item": line.get("item", line["description"]), "qty_per_unit": 1, "unit": line.get("unit"),
             "unit_cost": line["unit_cost"], "pack_size": line.get("pack_size")}
        if "waste_pct" in line:
            c["waste_pct"] = line["waste_pct"]
        comps.append((t, c, f"{where}.unit_cost"))
    if "hours" in line or "hours_per_unit" in line:
        c = {"role": line.get("role")}
        c.update({k: line[k] for k in ("hours", "hours_per_unit") if k in line})
        comps.append(("labor", c, f"{where}.hours"))
    if not comps:
        raise SpecError(f"{where}: line has no cost components")
    return comps


def price_line(line, cfg, where):
    for k in ("id", "description", "qty"):
        if k not in line:
            raise SpecError(f"{where}: missing {k}")
    qty = dec(line["qty"], f"{where}.qty")
    if qty <= 0:
        raise SpecError(f"{where}: qty must be > 0")
    at_cost_line = bool(line.get("at_cost"))
    out = {"id": str(line["id"]), "section": line.get("section", "Scope"), "description": line["description"],
           "qty": qty, "unit": line.get("unit", "ea"), "components": [], "notes": line.get("notes", "")}
    for ctype, c, cw in expand_components(line, where):
        row = {"type": ctype, "item": c.get("item") or c.get("role") or ctype, "hours": ZERO,
               "use_tax": ZERO, "waste": ZERO, "pack": None}
        if ctype == "labor":
            role = c.get("role")
            if role not in cfg.roles:
                raise SpecError(f"{cw}: role {role!r} not in labor.roles")
            wage, burden = cfg.roles[role]
            if "hours_per_unit" in c:
                hours = dec(c["hours_per_unit"], cw) * qty
            elif "hours" in c:
                hours = dec(c["hours"], cw)
            else:
                raise SpecError(f"{cw}: labor needs hours or hours_per_unit")
            rate = wage * (ONE + burden)
            row.update(item=c.get("item", role), hours=hours, order_qty=hours, unit="hr", wage=wage,
                       burden=burden, unit_cost=rate, cost=money(hours * rate))
        else:
            if "unit_cost" not in c:
                raise SpecError(f"{cw}: missing unit_cost")
            uc = dec(c["unit_cost"], f"{cw}.unit_cost")
            base = dec(c["qty_per_unit"], cw) * qty if "qty_per_unit" in c else dec(c.get("qty", 1), cw)
            waste = dec(c.get("waste_pct", 0), f"{cw}.waste_pct")
            order = base * (ONE + waste)
            pack = dec(c["pack_size"], f"{cw}.pack_size") if c.get("pack_size") else None
            if pack:
                order = (order / pack).to_integral_value(rounding=ROUND_CEILING) * pack
            cost = money(order * uc)
            row.update(base_qty=base, waste=waste, pack=pack, order_qty=order, unit=c.get("unit", "ea"),
                       waste_set="waste_pct" in c,
                       unit_cost=uc, cost=cost)
            if ctype == "materials" and cfg.tax_mode == "contractor_pays_on_cost":
                row["use_tax"] = money(cost * cfg.tax_rate)
        direct = row["cost"] + row["use_tax"]
        rule = ("at_cost", ZERO) if at_cost_line else cfg.pricing.rule(ctype)
        if rule[0] == "at_cost":
            cont = oh = ZERO
        else:
            cont = direct * cfg.cont
            oh = (direct + cont) * cfg.oh_pct + row["hours"] * cfg.oh_hr
        loaded = direct + cont + oh
        row.update(direct=direct, contingency=cont, overhead=oh, loaded=loaded, rule=rule,
                   sell=Pricing.sell(loaded, rule))
        out["components"].append(row)
    comps = out["components"]
    raw = sum((c["sell"] for c in comps), ZERO)
    if cfg.client_view == "lump_sum" or line.get("lump_sum"):
        out["unit_price"], out["amount"], out["lump"] = None, money(raw), True
    else:
        out["unit_price"] = money(raw / qty)
        out["amount"], out["lump"] = money(out["unit_price"] * qty), False
    taxable_raw = sum((c["sell"] for c in comps if c["type"] in cfg.taxable), ZERO)
    out["tax_share"] = (taxable_raw / raw).quantize(SHARE_Q, rounding=ROUND_HALF_UP) if (
        raw and cfg.tax_mode == "charge_on_sell") else ZERO
    out["taxable_amount"] = money(out["amount"] * out["tax_share"])
    out["sell_raw"] = raw
    for k in ("direct", "contingency", "overhead", "loaded", "hours", "use_tax"):
        out[k] = sum((c[k] for c in comps), ZERO)
    out["labor_cost"] = sum((c["cost"] for c in comps if c["type"] == "labor"), ZERO)
    out["loaded"] = money(out["loaded"])
    out["profit"] = out["amount"] - out["loaded"]
    return out


def total_block(lines, cfg):
    t = {k: sum((l[k] for l in lines), ZERO) for k in
         ("amount", "taxable_amount", "direct", "contingency", "overhead", "loaded", "hours", "labor_cost", "use_tax")}
    t["subtotal"] = t.pop("amount")
    t["tax"] = money(t["taxable_amount"] * cfg.tax_rate) if cfg.tax_mode == "charge_on_sell" else ZERO
    t["total"] = t["subtotal"] + t["tax"]
    t["profit"] = t["subtotal"] - t["loaded"]
    t["gross_profit"] = t["subtotal"] - t["direct"]
    s = t["subtotal"]
    t["net_margin"] = t["profit"] / s if s else ZERO
    t["gross_margin"] = t["gross_profit"] / s if s else ZERO
    t["markup_on_cost"] = t["profit"] / t["loaded"] if t["loaded"] else ZERO
    t["effective_hourly"] = (s - (t["direct"] - t["labor_cost"])) / t["hours"] if t["hours"] else None
    t["profit_per_hour"] = t["profit"] / t["hours"] if t["hours"] else None
    return t


# ---------------------------------------------------------------- checks
GENERAL_KW = {"permit (or say none required)": r"permit", "disposal / haul-away": r"dispos|haul|dumpster|debris",
              "mobilization / trip or travel": r"mobiliz|trip|travel|drive", "cleanup": r"clean"}
TRADE_KW = {
    "hvac": {"start-up / commissioning": r"start-?up|commission", "thermostat": r"thermostat",
             "line set": r"line ?set|refrigerant line", "condensate": r"condensate",
             "electrical disconnect / whip": r"disconnect|whip", "equipment pad or stand": r"\bpad\b|stand",
             "refrigerant recovery": r"recover", "plenum / transitions": r"plenum|transition"},
    "plumbing": {"shutoff valves": r"valve|shut-?off", "fittings": r"fitting", "access / patch": r"access|patch",
                 "code items (expansion tank, venting)": r"expansion|vent|code", "test / inspection": r"test|inspect"},
    "electrical": {"panel capacity / load calc": r"load|panel", "utility coordination": r"utility",
                   "patch / repair": r"patch|repair", "boxes and devices": r"\bbox|device", "inspection": r"inspect"},
    "roofing": {"tear-off": r"tear-?off", "decking allowance": r"deck|sheathing", "flashing": r"flashing",
                "drip edge": r"drip", "underlayment / ice barrier": r"underlayment|ice", "ventilation": r"vent"},
    "landscaping": {"utility locate": r"locate|811", "grading": r"grad", "irrigation": r"irrigat",
                    "delivery": r"deliver", "plant warranty": r"warrant"},
    "painting": {"surface prep": r"prep|patch|sand|caulk", "primer / spot prime": r"prime", "protection / masking":
                 r"mask|protect|drop", "furniture moving": r"furniture|move", "touch-up": r"touch"},
    "cleaning": {"supplies / consumables": r"suppl|consumable|chemical", "access / keys": r"access|key",
                 "frequency / visit count": r"weekly|monthly|visit|frequen"},
    "remodeling": {"demolition": r"demo", "protection": r"protect", "allowances": r"allowance",
                   "supervision / PM": r"supervis|project manag", "final clean": r"clean"},
    "general_contracting": {"supervision / PM": r"supervis|project manag", "general conditions": r"general cond|temp",
                            "insurance / bond": r"insur|bond", "allowances": r"allowance", "punch list": r"punch"},
    "it_av": {"cable pathways": r"cable|pathway|conduit", "configuration / programming": r"config|program",
              "testing / certification": r"test|certif", "labeling / documentation": r"label|document",
              "training / handover": r"train|handover", "mounting hardware": r"mount|bracket"},
}


def run_checks(spec, cfg, tiers, addons):
    w = []
    p = cfg.pricing
    w.append(("info", f"Pricing by {p.method} {pct(p.value, 2)} -> equivalent {'markup' if p.method == 'margin' else 'margin'} "
                      f"{pct(p.markup if p.method == 'margin' else p.margin, 2)}; price multiplier on loaded cost "
                      f"{(ONE + p.markup).quantize(D('0.0001'))}"))
    if p.margin > D("0.6"):
        w.append(("warn", f"Target margin {pct(p.margin)} is unusually high -- confirm this is not a markup number"))
    if cfg.burden == 0 and any(b == 0 for _, b in cfg.roles.values()):
        w.append(("warn", "Labor burden is 0% -- payroll taxes, workers comp, insurance and benefits are missing"))
    if cfg.oh_pct == 0 and cfg.oh_hr == 0:
        w.append(("warn", "No overhead recovery -- the target margin must then cover rent, trucks, office, insurance"))
    if cfg.tax_mode != "none" and cfg.tax_rate == 0:
        w.append(("warn", f"tax.mode is {cfg.tax_mode} but tax.rate is 0"))
    no_waste = []
    for tname, lines, tot in tiers:
        for l in lines:
            for c in l["components"]:
                if c["type"] == "materials" and not c.get("waste_set") and f"{l['id']}:{c['item']}" not in no_waste:
                    no_waste.append(f"{l['id']}:{c['item']}")
                if c["type"] != "labor" and c["unit_cost"] == 0:
                    w.append(("warn", f"{l['id']}: '{c['item']}' has a $0 unit cost"))
        eff_m = tot["net_margin"]
        if eff_m + D("0.005") < p.margin:
            w.append(("warn", f"{tname}: achieved net margin {pct(eff_m)} is below target {pct(p.margin)} "
                              f"(at-cost lines or lower sub/fee markups dilute it)"))
        if cfg.min_eff is not None and tot["effective_hourly"] is not None and tot["effective_hourly"] < cfg.min_eff:
            w.append(("warn", f"{tname}: effective hourly {fmt(tot['effective_hourly'])} is below the floor "
                              f"{fmt(cfg.min_eff)} -- raise price or cut hours"))
    if no_waste:
        w.append(("warn", f"{len(no_waste)} material components have no waste_pct (priced at 0% waste). Set it "
                          f"explicitly -- 0 for boxed equipment, 5-15% for cut or bulk goods: " + "; ".join(no_waste)))
    totals = [t[2]["subtotal"] for t in tiers]
    if len(totals) > 1 and totals != sorted(totals):
        w.append(("warn", "Option prices do not increase from first to last tier -- check tier order"))
    sched = (spec.get("payment_terms") or {}).get("schedule") or []
    if sched:
        s = sum((dec(m["pct"], "payment_terms.schedule.pct") for m in sched), ZERO)
        if s != ONE:
            raise SpecError(f"payment_terms.schedule percentages sum to {s}, must sum to 1")
    else:
        w.append(("warn", "No payment schedule -- add a deposit and milestones"))
    # scope completeness scan
    blob = " ".join([l["description"] + " " + " ".join(str(c["item"]) for c in l["components"])
                     for _, lines, _ in tiers for l in lines] +
                    [l["description"] for _, lines, _ in addons for l in lines] +
                    [str(x) for x in spec.get("exclusions", []) + spec.get("assumptions", [])]).lower()
    kws = dict(GENERAL_KW)
    kws.update(TRADE_KW.get(str(spec.get("trade", "")).lower().replace(" ", "_").replace("/", "_"), {}))
    for label, rx in kws.items():
        if not re.search(rx, blob):
            w.append(("check", f"Not found in scope, exclusions or assumptions: {label} -- price it or exclude it"))
    return w


# ---------------------------------------------------------------- build
def build(spec):
    cfg = Config(spec)
    base = spec.get("base_items") or []
    tier_specs = spec.get("tiers") or [{"name": "Base scope", "items": []}]
    if not base and not any(t.get("items") for t in tier_specs):
        raise SpecError("spec has no base_items and no tier items")
    ids = [str(i.get("id")) for i in base]
    tiers = []
    for ti, t in enumerate(tier_specs):
        items = list(base) + list(t.get("items") or [])
        tids = [str(i.get("id")) for i in items]
        dup = {x for x in tids if tids.count(x) > 1}
        if dup:
            raise SpecError(f"tier {t.get('name')}: duplicate line ids {sorted(dup)}")
        lines = [price_line(it, cfg, f"tiers[{ti}]({t.get('name')}).{it.get('id')}") for it in items]
        tiers.append((t.get("name", f"Option {ti + 1}"), lines, total_block(lines, cfg), t.get("description", "")))
    addons = []
    for ai, a in enumerate(spec.get("add_ons") or []):
        lines = [price_line(it, cfg, f"add_ons[{ai}].{it.get('id')}") for it in a.get("items") or []]
        addons.append((a.get("name", f"Add-on {ai + 1}"), lines, total_block(lines, cfg), a.get("description", "")))
    warnings = run_checks(spec, cfg, [t[:3] for t in tiers], [a[:3] for a in addons])
    return cfg, tiers, addons, warnings, ids


# ---------------------------------------------------------------- xlsx helpers
def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="C9CED6")
    return {
        "h": Font(bold=True, color="FFFFFF"), "hf": PatternFill("solid", fgColor="1F2A37"),
        "sec": Font(bold=True, color="1F2A37"), "secf": PatternFill("solid", fgColor="E8ECF1"),
        "b": Font(bold=True), "title": Font(bold=True, size=16, color="1F2A37"),
        "muted": Font(color="5B6472", size=9), "wrap": Alignment(wrap_text=True, vertical="top"),
        "border": Border(bottom=thin), "total": Font(bold=True, size=12),
        "money": '"$"#,##0.00', "pct": "0.00%", "qty": "#,##0.####",
    }


def _sheet_title(name, used):
    t = re.sub(r"[\[\]\*\?/\\:]", "", name)[:28] or "Sheet"
    base, n = t, 2
    while t in used:
        t = f"{base[:25]} {n}"
        n += 1
    used.add(t)
    return t


def write_client_xlsx(spec, cfg, tiers, addons, path):
    from openpyxl import Workbook
    s = _styles()
    wb = Workbook()
    used = set()
    biz, cli, q = spec.get("business") or {}, spec.get("client") or {}, spec.get("quote") or {}
    date, valid = quote_dates(q)
    taxed = cfg.tax_mode == "charge_on_sell"

    def header(ws, subtitle):
        ws["A1"] = biz.get("name", "Your Company")
        ws["A1"].font = s["title"]
        ws["A2"] = " | ".join(x for x in [biz.get("phone"), biz.get("email"), biz.get("license")] if x)
        ws["A2"].font = s["muted"]
        ws["A4"], ws["B4"] = "Quote", q.get("number", "Q-0001")
        ws["A5"], ws["B5"] = "Date", date.isoformat()
        ws["A6"], ws["B6"] = "Valid until", valid.isoformat()
        ws["D4"], ws["E4"] = "Prepared for", cli.get("name", "")
        ws["D5"], ws["E5"] = "Job site", cli.get("site_address", "")
        ws["A8"] = subtitle
        ws["A8"].font = s["sec"]
        for c in ("A4", "A5", "A6", "D4", "D5"):
            ws[c].font = s["b"]
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 58
        for col, wd in zip("CDEFG", (10, 8, 13, 14, 13)):
            ws.column_dimensions[col].width = wd

    def table(ws, lines, start, title):
        cols = ["Item", "Description", "Qty", "Unit", "Unit price", "Amount"] + (["Taxable"] if taxed else [])
        for i, h in enumerate(cols):
            c = ws.cell(row=start, column=i + 1, value=h)
            c.font, c.fill = s["h"], s["hf"]
        r = start + 1
        first = r
        section = None
        for l in lines:
            if l["section"] != section:
                section = l["section"]
                ws.cell(row=r, column=1, value=section).font = s["sec"]
                for col in range(1, len(cols) + 1):
                    ws.cell(row=r, column=col).fill = s["secf"]
                r += 1
            ws.cell(row=r, column=1, value=l["id"])
            ws.cell(row=r, column=2, value=l["description"]).alignment = s["wrap"]
            ws.cell(row=r, column=3, value=float(l["qty"])).number_format = s["qty"]
            ws.cell(row=r, column=4, value=l["unit"])
            if l["lump"]:
                ws.cell(row=r, column=5, value="incl.")
                ws.cell(row=r, column=6, value=float(l["amount"]))
            else:
                ws.cell(row=r, column=5, value=float(l["unit_price"])).number_format = s["money"]
                ws.cell(row=r, column=6, value=f"=ROUND(C{r}*E{r},2)")
            ws.cell(row=r, column=6).number_format = s["money"]
            if taxed:
                ws.cell(row=r, column=7, value=f"=ROUND(F{r}*{l['tax_share']},2)").number_format = s["money"]
            r += 1
        last = r - 1
        r += 1
        ws.cell(row=r, column=5, value="Subtotal").font = s["b"]
        ws.cell(row=r, column=6, value=f"=SUM(F{first}:F{last})").number_format = s["money"]
        sub = r
        if taxed:
            r += 1
            ws.cell(row=r, column=4, value=f"{cfg.tax_label} rate")
            ws.cell(row=r, column=5, value=float(cfg.tax_rate)).number_format = "0.000%"
            ws.cell(row=r, column=6, value=f"=ROUND(SUM(G{first}:G{last})*E{r},2)").number_format = s["money"]
            taxr = r
        r += 1
        ws.cell(row=r, column=5, value="Total").font = s["total"]
        ws.cell(row=r, column=6, value=f"=F{sub}+F{taxr}" if taxed else f"=F{sub}").number_format = s["money"]
        ws.cell(row=r, column=6).font = s["total"]
        if cfg.tax_mode == "contractor_pays_on_cost":
            ws.cell(row=r + 1, column=2, value="Prices include sales tax paid on materials.").font = s["muted"]
        return {"sub": f"F{sub}", "tax": f"F{taxr}" if taxed else None, "total": f"F{r}"}

    summary = wb.active
    summary.title = _sheet_title("Summary", used)
    refs = []
    for name, lines, tot, desc in tiers:
        ws = wb.create_sheet(_sheet_title(name, used))
        header(ws, f"{name}" + (f" -- {desc}" if desc else ""))
        refs.append((name, desc, ws.title, table(ws, lines, 10, name)))
    add_refs = []
    if addons:
        ws = wb.create_sheet(_sheet_title("Add-ons", used))
        header(ws, "Optional add-ons (priced individually, may be added to any option)")
        r = 10
        for name, lines, tot, desc in addons:
            ws.cell(row=r, column=1, value=name).font = s["b"]
            if desc:
                ws.cell(row=r, column=2, value=desc).font = s["muted"]
            ref = table(ws, lines, r + 1, name)
            add_refs.append((name, desc, ws.title, ref))
            r = int(ref["total"][1:]) + 3
    # summary sheet
    header(summary, spec.get("quote", {}).get("title", "Quote summary"))
    hdr = ["Option", "Description", "Subtotal", taxed and cfg.tax_label or "", "Total"]
    for i, h in enumerate(hdr):
        c = summary.cell(row=10, column=i + 1, value=h)
        c.font, c.fill = s["h"], s["hf"]
    r = 11
    for name, desc, st, ref in refs:
        summary.cell(row=r, column=1, value=name).font = s["b"]
        summary.cell(row=r, column=2, value=desc).alignment = s["wrap"]
        summary.cell(row=r, column=3, value=f"='{st}'!{ref['sub']}").number_format = s["money"]
        if ref["tax"]:
            summary.cell(row=r, column=4, value=f"='{st}'!{ref['tax']}").number_format = s["money"]
        summary.cell(row=r, column=5, value=f"='{st}'!{ref['total']}").number_format = s["money"]
        r += 1
    if add_refs:
        r += 1
        summary.cell(row=r, column=1, value="Optional add-ons").font = s["sec"]
        r += 1
        for name, desc, st, ref in add_refs:
            summary.cell(row=r, column=1, value=name)
            summary.cell(row=r, column=2, value=desc).alignment = s["wrap"]
            summary.cell(row=r, column=3, value=f"='{st}'!{ref['sub']}").number_format = s["money"]
            if ref["tax"]:
                summary.cell(row=r, column=4, value=f"='{st}'!{ref['tax']}").number_format = s["money"]
            summary.cell(row=r, column=5, value=f"='{st}'!{ref['total']}").number_format = s["money"]
            r += 1
    for col, wd in zip("ABCDE", (22, 58, 14, 13, 14)):
        summary.column_dimensions[col].width = wd
    # terms sheet
    ws = wb.create_sheet(_sheet_title("Terms", used))
    header(ws, "Payment schedule, exclusions, assumptions, acceptance")
    pt = spec.get("payment_terms") or {}
    sched = pt.get("schedule") or []
    r = 10
    if sched:
        cols = ["Milestone", "Share"] + [n for n, *_ in refs]
        for i, h in enumerate(cols):
            c = ws.cell(row=r, column=i + 1, value=h)
            c.font, c.fill = s["h"], s["hf"]
        first = r + 1
        for mi, m in enumerate(sched):
            rr = first + mi
            ws.cell(row=rr, column=1, value=m["milestone"])
            ws.cell(row=rr, column=2, value=float(dec(m["pct"], "pct"))).number_format = "0%"
            for ti, (_, _, st, ref) in enumerate(refs):
                col = 3 + ti
                L = ws.cell(row=rr, column=col).column_letter
                if mi < len(sched) - 1:
                    f = f"=ROUND('{st}'!{ref['total']}*B{rr},2)"
                else:  # last milestone takes the remainder so the schedule sums to the penny
                    f = f"='{st}'!{ref['total']}-SUM({L}{first}:{L}{rr - 1})" if mi else f"='{st}'!{ref['total']}"
                ws.cell(row=rr, column=col, value=f).number_format = s["money"]
        r = first + len(sched) + 1
    for label, key in (("Payment terms", "notes"),):
        for n in pt.get(key) or []:
            ws.cell(row=r, column=1, value=n).alignment = s["wrap"]
            r += 1
    for label, items in (("Exclusions", spec.get("exclusions")), ("Assumptions", spec.get("assumptions")),
                         ("Terms", spec.get("terms"))):
        if items:
            r += 1
            ws.cell(row=r, column=1, value=label).font = s["sec"]
            r += 1
            for it in items:
                ws.cell(row=r, column=1, value=f"- {it}").alignment = s["wrap"]
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
                ws.row_dimensions[r].height = 30 if len(str(it)) > 95 else 15
                r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"This quote is valid until {valid.isoformat()}. Work not listed is excluded "
                                   f"and will be quoted by written change order.").alignment = s["wrap"]
    r += 2
    for lab in ("Option selected", "Add-ons selected", "Client signature", "Printed name", "Date"):
        ws.cell(row=r, column=1, value=lab).font = s["b"]
        ws.cell(row=r, column=2).border = s["border"]
        r += 2
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    wb.save(path)


def write_internal_xlsx(spec, cfg, tiers, addons, warnings, path):
    from openpyxl import Workbook
    s = _styles()
    wb = Workbook()
    used = set()
    ws = wb.active
    ws.title = _sheet_title("Summary", used)
    ws["A1"] = "INTERNAL COST SHEET -- do not send to client"
    ws["A1"].font = s["title"]
    p = cfg.pricing
    info = [("Pricing method", p.method), ("Target margin", float(p.margin)), ("Equivalent markup", float(p.markup)),
            ("Default labor burden", float(cfg.burden)), ("Overhead % of cost", float(cfg.oh_pct)),
            ("Overhead per labor hour", float(cfg.oh_hr)), ("Contingency", float(cfg.cont)),
            ("Tax mode", cfg.tax_mode), ("Tax rate", float(cfg.tax_rate))]
    for i, (k, v) in enumerate(info):
        ws.cell(row=3 + i, column=1, value=k).font = s["b"]
        c = ws.cell(row=3 + i, column=2, value=v)
        if isinstance(v, float):
            c.number_format = s["money"] if "hour" in k else "0.00%"
    rows = [("Price (pre-tax)", "subtotal", "m"), ("Tax to client", "tax", "m"), ("Client total", "total", "m"),
            ("Direct cost", "direct", "m"), ("  of which labor", "labor_cost", "m"),
            ("  of which tax on materials", "use_tax", "m"), ("Contingency", "contingency", "m"),
            ("Overhead", "overhead", "m"), ("Loaded cost", "loaded", "m"), ("Gross profit (price - direct)",
            "gross_profit", "m"), ("Gross margin", "gross_margin", "p"), ("Net profit (price - loaded)", "profit", "m"),
            ("Net margin", "net_margin", "p"), ("Markup on loaded cost", "markup_on_cost", "p"),
            ("Labor hours", "hours", "q"), ("Effective hourly (price - non-labor direct) / hours",
            "effective_hourly", "m"), ("Net profit per labor hour", "profit_per_hour", "m")]
    top = 14
    groups = [(n, t) for n, _, t, _ in tiers] + [(f"Add-on: {n}", t) for n, _, t, _ in addons]
    ws.cell(row=top, column=1, value="Metric").font = s["h"]
    ws.cell(row=top, column=1).fill = s["hf"]
    for gi, (n, _) in enumerate(groups):
        c = ws.cell(row=top, column=2 + gi, value=n)
        c.font, c.fill = s["h"], s["hf"]
        ws.column_dimensions[c.column_letter].width = 18
    for ri, (label, key, kind) in enumerate(rows):
        ws.cell(row=top + 1 + ri, column=1, value=label)
        for gi, (_, t) in enumerate(groups):
            v = t.get(key)
            c = ws.cell(row=top + 1 + ri, column=2 + gi, value=None if v is None else float(v))
            c.number_format = {"m": s["money"], "p": s["pct"], "q": s["qty"]}[kind]
    r = top + len(rows) + 2
    ws.cell(row=r, column=1, value="Checks").font = s["sec"]
    for lvl, msg in warnings:
        r += 1
        ws.cell(row=r, column=1, value=f"[{lvl.upper()}] {msg}")
    ws.column_dimensions["A"].width = 52
    cols = ["Line", "Section", "Description", "Type", "Item", "Base qty", "Waste", "Order qty", "Unit", "Unit cost",
            "Hours", "Wage", "Burden", "Direct cost", "Tax on cost", "Contingency", "Overhead", "Loaded cost",
            "Pricing rule", "Sell (unrounded)"]
    for name, lines, tot, _ in tiers + addons:
        sh = wb.create_sheet(_sheet_title(f"Cost {name}", used))
        for i, h in enumerate(cols):
            c = sh.cell(row=1, column=i + 1, value=h)
            c.font, c.fill = s["h"], s["hf"]
        r = 2
        for l in lines:
            for c in l["components"]:
                vals = [l["id"], l["section"], l["description"], c["type"], c["item"],
                        float(c.get("base_qty", c["hours"])), float(c["waste"]), float(c["order_qty"]), c["unit"],
                        float(c["unit_cost"]), float(c["hours"]), float(c["wage"]) if "wage" in c else None,
                        float(c["burden"]) if "burden" in c else None, float(c["cost"]), float(c["use_tax"]),
                        float(c["contingency"]), float(c["overhead"]), float(c["loaded"]),
                        f"{c['rule'][0]} {pct(c['rule'][1], 2)}" if c["rule"][0] != "at_cost" else "at cost",
                        float(c["sell"])]
                for i, v in enumerate(vals):
                    cell = sh.cell(row=r, column=i + 1, value=v)
                    if i in (9, 11, 13, 14, 15, 16, 17, 19):
                        cell.number_format = s["money"]
                    elif i in (6, 12):
                        cell.number_format = "0.0%"
                r += 1
            sh.cell(row=r, column=3, value=f"{l['id']} client amount / profit / margin").font = s["b"]
            sh.cell(row=r, column=18, value=float(l["loaded"])).number_format = s["money"]
            sh.cell(row=r, column=20, value=float(l["amount"])).number_format = s["money"]
            sh.cell(row=r, column=21, value=float(l["profit"])).number_format = s["money"]
            sh.cell(row=r, column=22, value=f"=IF(T{r}=0,0,U{r}/T{r})").number_format = "0.0%"
            r += 2
        sh.cell(row=r, column=3, value="TOTAL direct (cost + tax on cost) / loaded / price").font = s["b"]
        sh.cell(row=r, column=14, value=f"=SUM(N2:N{r - 1})+SUM(O2:O{r - 1})").number_format = s["money"]
        sh.cell(row=r, column=18, value=float(tot["loaded"])).number_format = s["money"]
        sh.cell(row=r, column=20, value=float(tot["subtotal"])).number_format = s["money"]
        for col, wd in zip("ABCDE", (8, 16, 40, 10, 30)):
            sh.column_dimensions[col].width = wd
        sh.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------- pdf
def write_pdf(spec, cfg, tiers, addons, path):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        return False
    st = getSampleStyleSheet()
    biz, cli, q = spec.get("business") or {}, spec.get("client") or {}, spec.get("quote") or {}
    date, valid = quote_dates(q)
    ink = colors.HexColor("#1F2A37")
    el = [Paragraph(f"<b>{biz.get('name', 'Your Company')}</b>", st["Title"]),
          Paragraph(" | ".join(x for x in [biz.get("phone"), biz.get("email"), biz.get("license")] if x), st["Normal"]),
          Spacer(1, 8),
          Paragraph(f"<b>Quote {q.get('number', '')}</b> &nbsp; Date {date.isoformat()} &nbsp; Valid until "
                    f"{valid.isoformat()}", st["Normal"]),
          Paragraph(f"Prepared for: {cli.get('name', '')} &nbsp; Job site: {cli.get('site_address', '')}", st["Normal"]),
          Spacer(1, 6), Paragraph(q.get("title", ""), st["Heading2"])]
    taxed = cfg.tax_mode == "charge_on_sell"

    def tbl(lines, tot):
        data = [["Item", "Description", "Qty", "Unit", "Unit price", "Amount"]]
        for l in lines:
            data.append([l["id"], Paragraph(l["description"], st["BodyText"]), f"{l['qty'].normalize():f}", l["unit"],
                         "incl." if l["lump"] else fmt(l["unit_price"]), fmt(l["amount"])])
        data.append(["", "", "", "", "Subtotal", fmt(tot["subtotal"])])
        if taxed:
            data.append(["", "", "", "", cfg.tax_label, fmt(tot["tax"])])
        data.append(["", "", "", "", "Total", fmt(tot["total"])])
        t = Table(data, colWidths=[0.6 * inch, 3.4 * inch, 0.6 * inch, 0.5 * inch, 0.9 * inch, 1.0 * inch], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), ink), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
                               ("LINEBELOW", (0, 1), (-1, -4), 0.25, colors.HexColor("#C9CED6")),
                               ("FONTNAME", (4, -1), (-1, -1), "Helvetica-Bold")]))
        return t

    for name, lines, tot, desc in tiers:
        el += [Spacer(1, 10), Paragraph(f"<b>{name}</b>" + (f" -- {desc}" if desc else ""), st["Heading3"]),
               tbl(lines, tot)]
    if addons:
        el += [Spacer(1, 10), Paragraph("Optional add-ons", st["Heading3"])]
        for name, lines, tot, desc in addons:
            el += [Paragraph(f"<b>{name}</b> -- {fmt(tot['total'])}" + (f": {desc}" if desc else ""), st["BodyText"])]
    if cfg.tax_mode == "contractor_pays_on_cost":
        el.append(Paragraph("Prices include sales tax paid on materials.", st["Italic"]))
    pt = spec.get("payment_terms") or {}
    if pt.get("schedule"):
        el += [Spacer(1, 10), Paragraph("Payment schedule", st["Heading3"])]
        data = [["Milestone", "Share"] + [n for n, *_ in tiers]]
        sched = pt["schedule"]
        amounts = [schedule_amounts(t[2]["total"], sched) for t in tiers]
        for i, m in enumerate(sched):
            data.append([m["milestone"], pct(dec(m["pct"], "pct"), 0)] + [fmt(a[i]) for a in amounts])
        t = Table(data)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), ink), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
        el.append(t)
        for n in pt.get("notes") or []:
            el.append(Paragraph(n, st["BodyText"]))
    for label, items in (("Exclusions", spec.get("exclusions")), ("Assumptions", spec.get("assumptions")),
                         ("Terms", spec.get("terms"))):
        if items:
            el += [Spacer(1, 8), Paragraph(label, st["Heading3"])] + [Paragraph(f"- {x}", st["BodyText"]) for x in items]
    el += [Spacer(1, 14), Paragraph(f"Valid until {valid.isoformat()}. Work not listed is excluded and will be quoted by "
                                    f"written change order.", st["BodyText"]), Spacer(1, 18)]
    for lab in ("Option selected", "Client signature", "Date"):
        el += [Paragraph(f"{lab}: ________________________________________", st["BodyText"]), Spacer(1, 10)]
    SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                      topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=f"Quote {q.get('number', '')}").build(el)
    return True


# ---------------------------------------------------------------- misc
def quote_dates(q):
    raw = q.get("date")
    date = dt.date.fromisoformat(str(raw)) if raw else dt.date.today()
    return date, date + dt.timedelta(days=int(q.get("valid_days", 30)))


def schedule_amounts(total, sched):
    out = [money(total * dec(m["pct"], "pct")) for m in sched[:-1]]
    return out + [total - sum(out, ZERO)]


def explain(cfg, tiers, addons, line_id):
    for name, lines, _, _ in tiers + addons:
        for l in lines:
            if l["id"] == line_id:
                print(f"Line {line_id} ({name}): {l['description']} -- qty {l['qty']} {l['unit']}")
                for c in l["components"]:
                    if c["type"] == "labor":
                        print(f"  labor {c['item']}: {c['hours']} hr x ${c['wage']} x (1 + {c['burden']}) = "
                              f"{c['hours']} x ${c['unit_cost']} = {fmt(c['cost'])}")
                    else:
                        extra = f", pack {c['pack']} -> order {c['order_qty']}" if c["pack"] else f" = {c['order_qty']}"
                        print(f"  {c['type']} {c['item']}: {c['base_qty']} x (1 + {c['waste']}){extra} x "
                              f"${c['unit_cost']} = {fmt(c['cost'])}")
                    if c["use_tax"]:
                        print(f"    + tax paid on cost {fmt(c['use_tax'])}")
                    print(f"    contingency {fmt(c['contingency'])}, overhead {fmt(c['overhead'])}, loaded "
                          f"{c['loaded'].quantize(D('0.0001'))}; rule {c['rule'][0]} {c['rule'][1]} -> sell "
                          f"{c['sell'].quantize(D('0.0001'))}")
                if l["lump"]:
                    print(f"  line sell {l['sell_raw'].quantize(D('0.0001'))} -> amount {fmt(l['amount'])} (lump sum)")
                else:
                    print(f"  line sell {l['sell_raw'].quantize(D('0.0001'))} / qty {l['qty']} = unit price "
                          f"{fmt(l['unit_price'])}; amount = {l['qty']} x {fmt(l['unit_price'])} = {fmt(l['amount'])}")
                print(f"  loaded cost {fmt(l['loaded'])}, profit {fmt(l['profit'])}, line net margin "
                      f"{pct(l['profit'] / l['amount']) if l['amount'] else 'n/a'}; taxable share {l['tax_share']}")
                return
    sys.exit(f"line id {line_id!r} not found")


def to_jsonable(x):
    if isinstance(x, D):
        return float(x)
    if isinstance(x, dict):
        return {k: to_jsonable(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [to_jsonable(v) for v in x]
    return x


def main():
    ap = argparse.ArgumentParser(description="Build a trade quote from a YAML/JSON spec.")
    ap.add_argument("spec", nargs="?")
    ap.add_argument("--out", default=".")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--check", action="store_true", help="validate and print totals, write nothing")
    ap.add_argument("--explain", metavar="LINE_ID")
    ap.add_argument("--margin-to-markup", type=str)
    ap.add_argument("--markup-to-margin", type=str)
    a = ap.parse_args()
    try:
        if a.margin_to_markup:
            m = dec(a.margin_to_markup, "margin")
            print(f"margin {pct(m, 2)} = markup {pct(margin_to_markup(m), 2)}; price = cost / (1 - {m}) "
                  f"= cost x {(ONE / (ONE - m)).quantize(D('0.0001'))}")
            return
        if a.markup_to_margin:
            mu = dec(a.markup_to_margin, "markup")
            print(f"markup {pct(mu, 2)} = margin {pct(markup_to_margin(mu), 2)}; price = cost x {ONE + mu}")
            return
        if not a.spec:
            ap.error("spec is required")
        spec = load_spec(a.spec)
        cfg, tiers, addons, warnings, _ = build(spec)
    except SpecError as e:
        sys.exit(f"SPEC ERROR: {e}")
    if a.explain:
        explain(cfg, tiers, addons, a.explain)
        return
    q = spec.get("quote") or {}
    print(f"# {q.get('number', 'Q-0001')} -- {q.get('title', '')}")
    print("| Option | Price | Tax | Total | Direct | Loaded | Net profit | Net margin | Gross margin | Hours | Eff. hourly |")
    print("|---|---|---|---|---|---|---|---|---|---|---|")
    for name, _, t, _ in tiers + [(f"+ {n}", l, t, d) for n, l, t, d in addons]:
        eh = fmt(t["effective_hourly"]) if t["effective_hourly"] is not None else "n/a"
        print(f"| {name} | {fmt(t['subtotal'])} | {fmt(t['tax'])} | {fmt(t['total'])} | {fmt(t['direct'])} | "
              f"{fmt(t['loaded'])} | {fmt(t['profit'])} | {pct(t['net_margin'])} | {pct(t['gross_margin'])} | "
              f"{t['hours'].quantize(D('0.01'))} | {eh} |")
    print()
    for lvl, msg in warnings:
        print(f"[{lvl.upper()}] {msg}")
    if a.check:
        return
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9_.-]+", "-", str(q.get("number", "quote")))
    files = []
    try:
        import openpyxl  # noqa: F401
        write_client_xlsx(spec, cfg, tiers, addons, out / f"{stem}-client.xlsx")
        write_internal_xlsx(spec, cfg, tiers, addons, warnings, out / f"{stem}-internal.xlsx")
        files += [f"{stem}-client.xlsx", f"{stem}-internal.xlsx"]
    except ImportError:
        print("openpyxl not installed -- skipped xlsx (pip install openpyxl)")
    if not a.no_pdf:
        if write_pdf(spec, cfg, tiers, addons, out / f"{stem}-client.pdf"):
            files.append(f"{stem}-client.pdf")
        else:
            print("reportlab not installed -- skipped PDF (pip install reportlab)")
    summary = {"quote": q.get("number"), "options": {n: t for n, _, t, _ in tiers},
               "add_ons": {n: t for n, _, t, _ in addons},
               "lines": {n: {l["id"]: {k: l[k] for k in ("amount", "unit_price", "loaded", "profit", "hours")}
                             for l in lines} for n, lines, _, _ in tiers},
               "warnings": [f"[{lv}] {m}" for lv, m in warnings]}
    (out / f"{stem}-summary.json").write_text(json.dumps(to_jsonable(summary), indent=2))
    files.append(f"{stem}-summary.json")
    print("\nWrote: " + ", ".join(str(out / f) for f in files))


if __name__ == "__main__":
    main()
