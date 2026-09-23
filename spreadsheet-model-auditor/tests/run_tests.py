#!/usr/bin/env python3
"""Build fixtures, audit them, and check that every planted error is caught and
the clean model stays quiet. Uses LibreOffice for recalculation when present;
cache-dependent expectations are skipped (not failed) without it.

    python3 tests/run_tests.py
"""
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "scripts"))
sys.path.insert(0, HERE)

import audit_xlsx  # noqa: E402
import build_fixtures  # noqa: E402

CACHE_DEPENDENT = {("error_value", "Metrics", "B2"), ("stale_cached_value", "Commissions", "D8")}


def plant_stale_cache(path, sheet_title, cell, new_value):
    """Overwrite one cached <v> in the saved XML, as a non-recalculating tool would."""
    from openpyxl import load_workbook
    idx = load_workbook(path, read_only=True).sheetnames.index(sheet_title) + 1
    member = f"xl/worksheets/sheet{idx}.xml"
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == member:
                xml = data.decode("utf-8")
                xml, n = re.subn(rf'(<c r="{cell}"[^>]*>.*?<v>)([^<]*)(</v>)',
                                 rf"\g<1>{new_value}\g<3>", xml, count=1, flags=re.S)
                assert n == 1, f"could not plant stale value in {cell}"
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


def audit(path):
    return audit_xlsx.Auditor(path).run()


def main():
    work = tempfile.mkdtemp(prefix="sma-tests-")
    clean = os.path.join(work, "clean_model.xlsx")
    broken = os.path.join(work, "broken_model.xlsx")
    build_fixtures.build(clean, broken=False)
    build_fixtures.build(broken, broken=True)
    pricing = os.path.join(work, "clean_pricing.xlsx")
    build_fixtures.build_pricing(pricing)

    have_lo = bool(shutil.which("soffice") or shutil.which("libreoffice"))
    if have_lo:
        recalced = audit_xlsx.recalc_copy(broken)
        broken_calc = os.path.join(work, "broken_model_recalc.xlsx")
        shutil.copy(recalced, broken_calc)
        plant_stale_cache(broken_calc, "Commissions", "D8", 99999)
        clean_calc = audit_xlsx.recalc_copy(clean)
        pricing_calc = audit_xlsx.recalc_copy(pricing)
    else:
        broken_calc, clean_calc, pricing_calc = broken, clean, pricing
        print("LibreOffice not found: cache-dependent checks will be skipped.\n")

    failures = 0
    rep = audit(broken_calc)
    got = {(f["check"], f["sheet"], f["cell"]) for f in rep["findings"]}
    print(f"Broken model ({'recalculated' if have_lo else 'formulas only'}): "
          f"{len(rep['findings'])} findings {rep['summary']}")
    for exp in build_fixtures.EXPECTED:
        exp = tuple(exp)
        if exp in got:
            status = "CAUGHT"
        elif exp in CACHE_DEPENDENT and not have_lo:
            status = "SKIP  "
        else:
            status = "MISSED"
            failures += 1
        print(f"  {status}  {exp[0]:<24} {exp[1]}!{exp[2]}")

    raw = audit(broken)
    print(f"\nBroken model, raw (no cached values): {raw['summary']}")
    print("  no_cached_values notice present:",
          any(f["check"] == "no_cached_values" for f in raw["findings"]))

    for label, path in (("Clean operating model (raw)", clean),
                        ("Clean operating model (recalculated)", clean_calc),
                        ("Clean pricing + budget model (raw)", pricing),
                        ("Clean pricing + budget model (recalculated)", pricing_calc)):
        r = audit(path)
        noisy = [f for f in r["findings"] if f["severity"] not in ("info",)]
        print(f"\n{label}: {len(r['findings'])} findings {r['summary']}")
        for f in noisy:
            print(f"  FALSE POSITIVE  {f['severity']} {f['check']} {f['sheet']}!{f['cell']} "
                  f"{f['message']}")
        failures += len(noisy)

    out_md = os.path.join(work, "broken_model_report.md")
    with open(out_md, "w") as fh:
        fh.write(audit_xlsx.to_markdown(rep))
    print(f"\nSample report: {out_md}")
    print("RESULT:", "PASS" if not failures else f"FAIL ({failures})")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
