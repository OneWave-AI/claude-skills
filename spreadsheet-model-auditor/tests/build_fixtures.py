#!/usr/bin/env python3
"""Build test workbooks: a clean 12-month operating model and a broken copy with
planted errors. Writes clean_model.xlsx, broken_model.xlsx, and expected.json
(the planted errors the auditor must catch)."""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
FIRST, LAST = 2, 13  # month columns B..M
TOT = 14             # column N


def inputs_sheet(wb, growth):
    ws = wb.active
    ws.title = "Inputs"
    rows = [("Assumption", "Value"), ("Starting units", 1000), ("Monthly growth", growth),
            ("Price per unit", 49), ("COGS % of revenue", 0.35), ("Commission rate", 0.08),
            ("Rent", 4000), ("Payroll", 30000), ("Software", 1200)]
    for r, (a, b) in enumerate(rows, 1):
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
    for r in (3, 5, 6):
        ws.cell(r, 2).number_format = "0.0%"
    return ws


def forecast_sheet(wb, broken):
    ws = wb.create_sheet("Forecast")
    ws.cell(1, 1, "Line")
    for i, m in enumerate(MONTHS):
        ws.cell(1, FIRST + i, m)
    ws.cell(1, TOT, "Total")
    labels = {2: "Units", 3: "Price", 4: "Revenue", 5: "COGS", 6: "Gross profit",
              7: "Rent", 8: "Payroll", 9: "Commission", 10: "Software",
              11: "Total opex", 12: "Operating income"}
    for r, lab in labels.items():
        ws.cell(r, 1, lab)
    for c in range(FIRST, LAST + 1):
        col, prev = L(c), L(c - 1)
        ws[f"{col}2"] = "=Inputs!$B$2" if c == FIRST else f"={prev}2*(1+Inputs!$B$3)"
        ws[f"{col}3"] = "=Inputs!$B$4"
        ws[f"{col}4"] = f"={col}2*{col}3"
        ws[f"{col}5"] = f"={col}4*Inputs!$B$5"
        ws[f"{col}6"] = f"={col}4-{col}5"
        ws[f"{col}7"] = "=Inputs!$B$7"
        ws[f"{col}8"] = "=Inputs!$B$8"
        ws[f"{col}9"] = f"={col}4*Inputs!$B$6"
        ws[f"{col}10"] = "=Inputs!$B$9"
        # Clean: total covers rows 7-10. Broken: the Software row was inserted
        # below the old range and the SUM never extended.
        ws[f"{col}11"] = f"=SUM({col}7:{col}9)" if broken else f"=SUM({col}7:{col}10)"
        ws[f"{col}12"] = f"={col}6-{col}11"
    for r in range(2, 13):
        fn = "AVERAGE" if r == 3 else "SUM"
        ws.cell(r, TOT, f"={fn}({L(FIRST)}{r}:{L(LAST)}{r})")
    if broken:
        ws["F5"] = "=F4*0.35"                     # hardcoded rate + inconsistent
        ws["H4"] = 52000                          # pasted over formula
        ws["J9"] = "=J4*Inputs!$B$7"              # wrong input row, inconsistent
        # Marketing row: one month entered in thousands instead of dollars
        ws.cell(14, 1, "Marketing spend")
        for c in range(FIRST, LAST + 1):
            ws.cell(14, c, 2 if c == 7 else 2000)
        # Refunds with mixed signs
        ws.cell(15, 1, "Refunds")
        for c in range(FIRST, LAST + 1):
            ws.cell(15, c, 300 if c in (4, 9) else -300)
        # Hidden manual-adjustment row feeding operating income
        ws.cell(16, 1, "Manual adj")
        ws.cell(16, 5, 7500)
        ws["E12"] = "=E6-E11+E16"
        ws.row_dimensions[16].hidden = True
    return ws


def commissions_sheet(wb, broken):
    ws = wb.create_sheet("Commissions")
    for c, h in enumerate(["Rep", "Bookings", "Rate", "Commission"], 1):
        ws.cell(1, c, h)
    bookings = [42000, 38500, 51000, 27000, 61000, 33500]
    for i, b in enumerate(bookings):
        r = 2 + i
        ws.cell(r, 1, f"Rep {chr(65 + i)}")
        ws.cell(r, 2, b)
        ws.cell(r, 3, "=Inputs!$B$6")
        ws.cell(r, 4, f"=B{r}*C{r}")
    ws.cell(8, 1, "Total")
    ws.cell(8, 2, 250000 if broken else "=SUM(B2:B7)")  # parts sum to 253000
    ws.cell(8, 4, "=SUM(D2:D7)")
    return ws


def metrics_sheet(wb):
    ws = wb.create_sheet("Metrics")
    rows = [
        ("Metric", "Value"),
        ("Revenue per customer", "=Forecast!N4/B3"),     # B3 blank -> #DIV/0!
        ("Customers", None),
        ("Broken link", "=#REF!*2"),                   # deleted reference
        ("Commission check", "=SUM(Commissions!D:D)"),  # whole column incl. total row
        ("Base amount", 1000),
        ("Processing fee", "=B8*0.02"),                # circular with B8
        ("Amount incl fee", "=B6+B7"),
        ("", None),
        ("Budget link", "='[Budget FY25.xlsx]Inputs'!B2"),
        ("Dynamic lookup", '=INDIRECT("Forecast!B4")'),
        ("As of", "=TODAY()"),
    ]
    for r, (a, b) in enumerate(rows, 1):
        if a:
            ws.cell(r, 1, a)
        if b is not None:
            ws.cell(r, 2, b)
    return ws


def build(path, broken):
    wb = Workbook()
    inputs_sheet(wb, 3 if broken else 0.03)  # broken: 3 typed meaning 3%
    forecast_sheet(wb, broken)
    commissions_sheet(wb, broken)
    if broken:
        metrics_sheet(wb)
        s = wb.create_sheet("Scratch")
        s["A1"] = "old plug"
        s["B1"] = "=Forecast!N12*0.1"
        s.sheet_state = "hidden"
    wb.save(path)


def build_pricing(path):
    """Second clean model with a different shape: tier lookup, quote lines,
    a departmental budget with spaced subtotal blocks, and a named range."""
    from openpyxl.workbook.defined_name import DefinedName
    wb = Workbook()
    t = wb.active
    t.title = "Tiers"
    for r, row in enumerate([("Min qty", "Unit price", "Discount"), (0, 25, 0),
                             (50, 22, 0.05), (200, 19, 0.1), (1000, 16, 0.15)], 1):
        for c, v in enumerate(row, 1):
            t.cell(r, c, v)
    t["E1"], t["F1"] = "Tax rate", 0.0725
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="Tiers!$F$1")
    q = wb.create_sheet("Quote")
    for c, h in enumerate(["Item", "Qty", "Unit price", "Discount", "Line total", "Tax"], 1):
        q.cell(1, c, h)
    for i, (item, qty) in enumerate([("Widget", 40), ("Gadget", 250), ("Bracket", 1200),
                                     ("Cable", 75), ("Mount", 10)]):
        r = 2 + i
        q.cell(r, 1, item)
        q.cell(r, 2, qty)
        q.cell(r, 3, f"=VLOOKUP(B{r},Tiers!$A$2:$C$5,2,TRUE)")
        q.cell(r, 4, f"=VLOOKUP(B{r},Tiers!$A$2:$C$5,3,TRUE)")
        q.cell(r, 5, f"=ROUND(B{r}*C{r}*(1-D{r}),2)")
        q.cell(r, 6, f"=ROUND(E{r}*TaxRate,2)")
    q["A7"], q["E7"], q["F7"] = "Total", "=SUM(E2:E6)", "=SUM(F2:F6)"
    q["A8"], q["E8"] = "Avg price per unit", "=IFERROR(E7/SUM(B2:B6),0)"
    b = wb.create_sheet("Budget")
    b["A1"], b["B1"], b["C1"], b["D1"] = "Line", "Q1", "Q2", "FY"
    layout = [(2, "Sales salaries", 120000), (3, "Sales travel", 18000), (4, "Sales tools", 9000),
              (5, "Sales subtotal", None), (7, "Ops salaries", 90000), (8, "Ops software", 14000),
              (9, "Ops subtotal", None), (11, "Total spend", None)]
    for r, lab, v in layout:
        b.cell(r, 1, lab)
        for c in (2, 3):
            col = "B" if c == 2 else "C"
            if v is not None:
                b.cell(r, c, v if c == 2 else round(v * 1.04))
            elif r == 5:
                b.cell(r, c, f"=SUM({col}2:{col}4)")
            elif r == 9:
                b.cell(r, c, f"=SUM({col}7:{col}8)")
            else:
                b.cell(r, c, f"={col}5+{col}9")
        b.cell(r, 4, f"=SUM(B{r}:C{r})")
    wb.save(path)


EXPECTED = [
    ("hardcoded_in_formula", "Forecast", "F5"),
    ("inconsistent_formula", "Forecast", "F5"),
    ("pasted_over_formula", "Forecast", "H4"),
    ("inconsistent_formula", "Forecast", "J9"),
    ("inconsistent_formula", "Forecast", "E12"),
    ("sum_range_omission", "Forecast", "B11"),
    ("sum_range_omission", "Forecast", "M11"),
    ("unit_outlier", "Forecast", "G14"),
    ("mixed_sign_convention", "Forecast", "row 15"),
    ("hidden_content", "Forecast", "row 16"),
    ("total_does_not_foot", "Commissions", "B8"),
    ("percent_as_whole_number", "Inputs", "B3"),
    ("ref_to_empty", "Metrics", "B2"),
    ("error_value", "Metrics", "B2"),          # #DIV/0! (needs cached values)
    ("error_value", "Metrics", "B4"),          # #REF!
    ("whole_column_range", "Metrics", "B5"),
    ("double_count", "Metrics", "B5"),
    ("circular_reference", "Metrics", "B7"),
    ("external_link", "Metrics", "B10"),
    ("volatile_function", "Metrics", "B11"),
    ("volatile_function", "Metrics", "B12"),
    ("hidden_content", "Scratch", "-"),
    ("stale_cached_value", "Commissions", "D8"),  # planted by run_tests.py after recalc
]

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
    os.makedirs(out, exist_ok=True)
    build(os.path.join(out, "clean_model.xlsx"), broken=False)
    build(os.path.join(out, "broken_model.xlsx"), broken=True)
    build_pricing(os.path.join(out, "clean_pricing.xlsx"))
    with open(os.path.join(out, "expected.json"), "w") as fh:
        json.dump(EXPECTED, fh, indent=1)
    print(f"wrote fixtures to {out}")
