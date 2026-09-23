#!/usr/bin/env python3
"""Structural integrity audit for business spreadsheet models (.xlsx / .xlsm).

Reads the workbook twice with openpyxl -- once for formulas, once for the
cached values Excel stored on last save -- and reports structural problems
with sheet, cell address, severity, why it matters, and the fix.

Usage:
    python3 audit_xlsx.py model.xlsx [--json out.json] [--md out.md]
                          [--recalc] [--max-per-check 25]

--recalc  Run LibreOffice headless to recalculate a copy first, so files
          written by scripts (no cached values) get values. Optional.

Exit code is 0 unless the file cannot be read. Requires: openpyxl.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict

try:
    from openpyxl import load_workbook
    from openpyxl.formula.tokenizer import Tokenizer, Token
    from openpyxl.utils import get_column_letter, column_index_from_string
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

SEVERITY_ORDER = ["critical", "high", "medium", "low", "info"]
ERROR_VALUES = {"#REF!", "#DIV/0!", "#N/A", "#VALUE!", "#NAME?", "#NUM!", "#NULL!",
                "#SPILL!", "#CALC!", "#GETTING_DATA"}
VOLATILE = {"INDIRECT", "OFFSET", "NOW", "TODAY", "RAND", "RANDBETWEEN", "RANDARRAY",
            "INFO", "CELL"}
AGGREGATES = {"SUM", "SUBTOTAL", "AGGREGATE", "AVERAGE", "MIN", "MAX", "COUNT",
              "COUNTA", "SUMIF", "SUMIFS", "SUMPRODUCT"}
# Functions whose numeric arguments are structural (column index, digits, date parts).
STRUCTURAL_ARG_FUNCS = {"ROUND", "ROUNDUP", "ROUNDDOWN", "MROUND", "VLOOKUP", "HLOOKUP",
                        "XLOOKUP", "INDEX", "MATCH", "XMATCH", "OFFSET", "LEFT", "RIGHT",
                        "MID", "CHOOSE", "DATE", "DATEVALUE", "EOMONTH", "EDATE", "WEEKDAY",
                        "WORKDAY", "NETWORKDAYS", "TEXT", "LARGE", "SMALL", "RANK", "RANK.EQ",
                        "PERCENTILE", "QUARTILE", "SUBTOTAL", "AGGREGATE", "FIXED", "YEARFRAC",
                        "DATEDIF", "TRUNC", "INT", "SEQUENCE", "TAKE", "DROP", "SORT", "FILTER",
                        "CEILING", "FLOOR", "CEILING.MATH", "FLOOR.MATH", "ADDRESS", "FIND",
                        "SEARCH", "SUBSTITUTE", "REPT", "TIME", "COLUMN", "ROW", "IFERROR"}
UNIT_CONSTANTS = {12, 52, 100, 1000, 365, 24, 60, 4, 7, 1e6, 1000000, 3600}
TRIVIAL_CONSTANTS = {0, 1}
BLANK_AWARE_FUNCS = {"ISBLANK", "IF", "IFS", "IFERROR", "IFNA", "COUNTA", "COUNTBLANK", "N",
                     "ISNUMBER", "ISTEXT", "LEN"}
EXPENSE_WORDS = re.compile(r"\b(cost|costs|expense|expenses|cogs|opex|refund|refunds|"
                           r"discount|discounts|payroll|rent|salar(y|ies)|fees?|"
                           r"returns|chargebacks?)\b", re.I)
# Google Sheets live-data functions; a Sheets .xlsx export keeps them as formula text
# (often wrapped in __xludf.DUMMYFUNCTION) with the last computed value cached.
SHEETS_LIVE = re.compile(r"\b(IMPORTRANGE|IMPORTDATA|IMPORTHTML|IMPORTXML|IMPORTFEED|"
                         r"GOOGLEFINANCE)\s*\(", re.I)
TOTAL_WORDS = re.compile(r"\b(total|subtotal|sum|grand total)\b", re.I)

CELL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
COL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})$")
ROW_RE = re.compile(r"^(\$?)(\d+)$")
MAX_EXPAND = 200_000

GUIDE = {
    "hardcoded_in_formula": (
        "A literal number buried in a formula is an assumption nobody can see, change in one "
        "place, or sensitivity-test. When the rate changes, some copies get updated and some do not.",
        "Move the number to a labeled input cell (or named range) and reference it."),
    "pasted_over_formula": (
        "A typed value sitting in a row or column of formulas stops updating. The model looks "
        "live but this cell is frozen at whatever someone pasted.",
        "Restore the formula from its neighbor (copy the adjacent cell across). If the override "
        "is intentional, move it to an input row and label it as an override."),
    "inconsistent_formula": (
        "The cell does not follow the pattern of its row or column. Usually a copy-paste slip, "
        "a range that did not extend, or a one-off tweak that silently changes one period.",
        "Compare with the dominant formula shown; re-fill the row/column from a correct cell. "
        "If the difference is intentional, document it in a comment or note column."),
    "error_value": (
        "An error value poisons every formula downstream and often hides behind IFERROR "
        "wrappers elsewhere. #REF! means a reference was deleted and cannot be recovered "
        "automatically.",
        "Fix the root cell (listed first); propagated errors clear once the root is fixed. For "
        "#REF!, rebuild the reference; for #DIV/0!, guard the denominator explicitly."),
    "ref_to_empty": (
        "A formula points at a blank cell, which Excel treats as zero. Often a reference that "
        "slipped one row/column when rows were inserted.",
        "Point the reference at the intended input, or delete it if it is dead."),
    "whole_column_range": (
        "Whole-column ranges pick up anything later typed in the column -- totals, notes, "
        "checks -- and silently double count. They also slow recalculation.",
        "Use an explicit bounded range or an Excel Table structured reference."),
    "double_count": (
        "The SUM range includes a subtotal that already sums cells in the same range, so "
        "those items are counted twice.",
        "Exclude the subtotal rows from the range, or sum only the subtotals."),
    "circular_reference": (
        "The cells depend on themselves. Excel either refuses to calculate or, with iterative "
        "calculation on, converges to a number that depends on the starting value.",
        "Break the loop: usually interest-on-average-balance or a fee on a total that includes "
        "the fee. Use a prior-period balance, a closed-form solution, or an explicit switch cell."),
    "sum_range_omission": (
        "Numeric rows sit right next to the summed range but are not in it. Classic cause: a "
        "row inserted at the edge of a range, which Excel does not always auto-extend.",
        "Extend the range to cover the omitted rows, or insert rows inside the range in future."),
    "total_does_not_foot": (
        "The total does not equal the sum of its parts. Anyone reading the total is reading a "
        "different number from the detail.",
        "Replace the typed total with a SUM of the detail, or find the missing/changed line."),
    "hardcoded_total": (
        "The total is typed, not calculated. It matches today but will not move when any "
        "line above changes.",
        "Replace it with =SUM(...) over the detail rows."),
    "stale_cached_value": (
        "The value stored in the file differs from what its formula computes from the stored "
        "inputs. The file was likely saved in manual-calculation mode or edited by a tool that "
        "does not recalculate, so what people see is not what the model says.",
        "Recalculate (F9 in Excel, or --recalc here) and re-check; switch calculation to automatic."),
    "external_link": (
        "The number depends on another file that may have moved, changed, or never be sent "
        "with this one. Recipients see stale values or a broken-link prompt.",
        "Paste the linked inputs into an Inputs sheet with a source/date note, or document the "
        "link and its owner."),
    "cross_sheet_link": (
        "Cross-sheet references are normal, but a sheet with many inbound links is a single "
        "point of failure, and references into a sheet usually break when it is restructured.",
        "No action needed unless the map looks surprising; keep inputs on one sheet."),
    "volatile_function": (
        "Volatile functions recalculate on every change. INDIRECT and OFFSET hide their "
        "precedents from auditing tools (including this one); NOW/TODAY/RAND make the model "
        "give different answers on different days.",
        "Replace INDIRECT/OFFSET with INDEX or direct references; put TODAY() in one labeled "
        "as-of cell; remove RAND from anything that is reported."),
    "hidden_content": (
        "Hidden sheets, rows, and columns that contain formulas or inputs are where stale "
        "overrides and forgotten plugs live. Reviewers cannot check what they cannot see.",
        "Unhide and review. Delete what is dead; move live inputs to a visible Inputs area."),
    "percent_as_whole_number": (
        "A percent-formatted cell holding a value above 1 displays as hundreds of percent "
        "(7 shows as 700%). Usually someone typed 7 meaning 7%.",
        "Enter the rate as a decimal (0.07) or confirm the intent."),
    "unit_outlier": (
        "One input is ~100x or ~1000x its neighbors. Common when one period is entered in "
        "dollars and the rest in thousands, or a percent as a whole number.",
        "Confirm the unit of every input in the row; label units in the row header."),
    "mixed_sign_convention": (
        "Some expense lines are entered positive and others negative. Totals that add them "
        "will partially net instead of summing.",
        "Pick one convention (costs positive and subtracted, or costs negative and added) and "
        "apply it to every line."),
    "no_cached_values": (
        "The file has formulas but no stored results -- it was written by a script or a tool "
        "that does not calculate. Value-based checks (errors, footing, stale cache) were skipped.",
        "Re-run with --recalc (needs LibreOffice), or open and save the file in Excel first."),
}


# --------------------------------------------------------------------------- helpers

def formula_text(value):
    if value is None:
        return None
    if isinstance(value, str):
        return value if value.startswith("=") else None
    text = getattr(value, "text", None)  # ArrayFormula
    if isinstance(text, str):
        return text if text.startswith("=") else "=" + text
    return None


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def addr(col, row):
    return f"{get_column_letter(col)}{row}"


def split_sheet(operand):
    """Return (sheet_or_None, a1part, external_bool)."""
    if "!" not in operand:
        ext = operand.startswith("[")
        return None, operand, ext
    idx = operand.rfind("!")
    sheet, a1 = operand[:idx], operand[idx + 1:]
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    external = "[" in sheet or "]" in sheet
    return sheet, a1, external


def parse_ref(operand, cur_sheet, sheetnames):
    """Parse a RANGE operand into a reference dict, or None for names/unknown."""
    sheet, a1, external = split_sheet(operand)
    ref = {"sheet": sheet or cur_sheet, "external": external, "raw": operand}
    if external:
        ref["kind"] = "external"
        return ref
    if ref["sheet"] not in sheetnames:
        ref["kind"] = "unknown"
        return ref
    parts = a1.split(":")
    if len(parts) == 1:
        m = CELL_RE.match(parts[0])
        if not m:
            ref["kind"] = "name"
            ref["name"] = parts[0]
            return ref
        c, r = column_index_from_string(m.group(2).upper()), int(m.group(4))
        ref.update(kind="cell", min_col=c, max_col=c, min_row=r, max_row=r,
                   abs=[(m.group(1) == "$", m.group(3) == "$")])
        return ref
    if len(parts) == 2:
        m1, m2 = CELL_RE.match(parts[0]), CELL_RE.match(parts[1])
        if m1 and m2:
            c1, r1 = column_index_from_string(m1.group(2).upper()), int(m1.group(4))
            c2, r2 = column_index_from_string(m2.group(2).upper()), int(m2.group(4))
            ref.update(kind="range", min_col=min(c1, c2), max_col=max(c1, c2),
                       min_row=min(r1, r2), max_row=max(r1, r2))
            return ref
        k1, k2 = COL_RE.match(parts[0]), COL_RE.match(parts[1])
        if k1 and k2:
            c1 = column_index_from_string(k1.group(2).upper())
            c2 = column_index_from_string(k2.group(2).upper())
            ref.update(kind="cols", min_col=min(c1, c2), max_col=max(c1, c2))
            return ref
        w1, w2 = ROW_RE.match(parts[0]), ROW_RE.match(parts[1])
        if w1 and w2:
            ref.update(kind="rows", min_row=min(int(w1.group(2)), int(w2.group(2))),
                       max_row=max(int(w1.group(2)), int(w2.group(2))))
            return ref
    ref["kind"] = "unknown"
    return ref


def to_r1c1_corner(token, row, col):
    m = CELL_RE.match(token)
    if m:
        c, r = column_index_from_string(m.group(2).upper()), int(m.group(4))
        cs = f"C{c}" if m.group(1) else f"C[{c - col}]"
        rs = f"R{r}" if m.group(3) else f"R[{r - row}]"
        return rs + cs
    m = COL_RE.match(token)
    if m:
        c = column_index_from_string(m.group(2).upper())
        return f"C{c}" if m.group(1) else f"C[{c - col}]"
    m = ROW_RE.match(token)
    if m:
        r = int(m.group(2))
        return f"R{r}" if m.group(1) else f"R[{r - row}]"
    return token.upper()


def r1c1(tokens, row, col):
    out = []
    for t in tokens:
        if t.type == Token.WSPACE:
            continue
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            sheet, a1, _ = split_sheet(t.value)
            conv = ":".join(to_r1c1_corner(p, row, col) for p in a1.split(":"))
            out.append((sheet.upper() + "!" if sheet else "") + conv)
        else:
            out.append(t.value.upper())
    return "".join(out)


# --------------------------------------------------------------------------- model

class Cell:
    __slots__ = ("sheet", "row", "col", "formula", "value", "cached", "tokens", "refs",
                 "funcs", "numbers", "r1c1", "hidden_row", "hidden_col", "number_format")

    def __init__(self, sheet, row, col):
        self.sheet, self.row, self.col = sheet, row, col
        self.formula = None
        self.value = None
        self.cached = None
        self.tokens = []
        self.refs = []
        self.funcs = []
        self.numbers = []
        self.r1c1 = None
        self.number_format = "General"

    @property
    def key(self):
        return (self.sheet, self.row, self.col)

    @property
    def a1(self):
        return addr(self.col, self.row)

    def display_value(self):
        return self.cached if self.formula else self.value


class Auditor:
    def __init__(self, path, max_per_check=25):
        self.path = path
        self.max_per_check = max_per_check
        self.findings = []
        self.cells = {}                       # (sheet,row,col) -> Cell
        self.by_sheet = defaultdict(dict)     # sheet -> {(row,col): Cell}
        self.formula_cells = defaultdict(dict)
        self.dims = {}
        self.hidden_rows = defaultdict(set)
        self.hidden_cols = defaultdict(set)
        self.sheet_state = {}
        self.names = {}
        self.has_cache = False
        self.precedents = {}

    # ---------------------------------------------------------------- loading
    def load(self):
        wb_f = load_workbook(self.path, data_only=False)
        wb_v = load_workbook(self.path, data_only=True)
        self.sheetnames = list(wb_f.sheetnames)
        self.ext_link_count = len(getattr(wb_f, "_external_links", []) or [])
        try:
            for name, dn in wb_f.defined_names.items():
                self.names[name.upper()] = list(dn.destinations)
        except Exception:
            pass
        for ws in wb_f.worksheets:
            wsv = wb_v[ws.title]
            self.sheet_state[ws.title] = ws.sheet_state
            self.dims[ws.title] = (ws.max_row, ws.max_column)
            try:
                for name, dn in ws.defined_names.items():
                    self.names.setdefault(name.upper(), list(dn.destinations))
            except Exception:
                pass
            for r, d in ws.row_dimensions.items():
                if d.hidden:
                    self.hidden_rows[ws.title].add(r)
            for _, d in ws.column_dimensions.items():
                if d.hidden and d.min and d.max:
                    self.hidden_cols[ws.title].update(range(d.min, d.max + 1))
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    cell = Cell(ws.title, c.row, c.column)
                    cell.number_format = c.number_format or "General"
                    f = formula_text(c.value)
                    if f:
                        cell.formula = f
                        cell.cached = wsv.cell(row=c.row, column=c.column).value
                        if cell.cached is not None:
                            self.has_cache = True
                        self._parse(cell)
                        self.formula_cells[ws.title][(c.row, c.column)] = cell
                    else:
                        cell.value = c.value
                    self.cells[cell.key] = cell
                    self.by_sheet[ws.title][(c.row, c.column)] = cell

    def _parse(self, cell):
        try:
            tokens = Tokenizer(cell.formula).items
        except Exception:
            return
        cell.tokens = tokens
        stack = []
        for t in tokens:
            if t.type == Token.FUNC and t.subtype == Token.OPEN:
                fn = re.sub(r"^(_XLFN\.|_XLWS\.)+", "", t.value[:-1].upper())
                stack.append(fn)
                cell.funcs.append(fn)
            elif t.type == Token.PAREN and t.subtype == Token.OPEN:
                stack.append(None)
            elif t.subtype == Token.CLOSE and stack:
                stack.pop()
            elif t.type == Token.OPERAND and t.subtype == Token.RANGE:
                ref = parse_ref(t.value, cell.sheet, self.sheetnames)
                ref["context"] = [s for s in stack if s]
                cell.refs.append(ref)
            elif t.type == Token.OPERAND and t.subtype == Token.NUMBER:
                try:
                    num = float(t.value)
                except ValueError:
                    continue
                cell.numbers.append((num, [s for s in stack if s]))
        cell.r1c1 = r1c1(tokens, cell.row, cell.col)

    # ---------------------------------------------------------------- utils
    def add(self, check, severity, sheet, cell, message, formula=None, value=None, extra=None):
        why, fix = GUIDE.get(check, ("", ""))
        f = {"check": check, "severity": severity, "sheet": sheet, "cell": cell,
             "message": message, "why": why, "fix": fix}
        if formula is not None:
            f["formula"] = formula
        if value is not None:
            f["value"] = value if isinstance(value, (int, float, str, bool)) else str(value)
        if extra:
            f.update(extra)
        self.findings.append(f)

    def get(self, sheet, row, col):
        return self.by_sheet[sheet].get((row, col))

    def numeric_value(self, cell):
        if cell is None:
            return None
        v = cell.display_value()
        return v if is_number(v) else None

    def label_for_row(self, sheet, row):
        for c in range(1, 4):
            cell = self.get(sheet, row, c)
            if cell and isinstance(cell.value, str) and cell.value.strip():
                return cell.value.strip()
        return ""

    def label_for_col(self, sheet, row, col):
        for r in range(row - 1, 0, -1):
            cell = self.get(sheet, r, col)
            if cell and isinstance(cell.value, str) and cell.value.strip():
                return cell.value.strip()
        return ""

    def expand(self, ref):
        """Yield (sheet,row,col) cells a ref covers, bounded by the sheet's used range."""
        kind = ref.get("kind")
        sheet = ref.get("sheet")
        if kind == "name":
            for dest_sheet, rng in self.names.get(ref["name"].upper(), []):
                sub = parse_ref(f"'{dest_sheet}'!{rng}", sheet, self.sheetnames)
                yield from self.expand(sub)
            return
        if kind not in ("cell", "range", "cols", "rows"):
            return
        max_r, max_c = self.dims.get(sheet, (0, 0))
        r1, r2 = ref.get("min_row", 1), min(ref.get("max_row", max_r), max_r)
        c1, c2 = ref.get("min_col", 1), min(ref.get("max_col", max_c), max_c)
        if kind == "cell":
            yield (sheet, ref["min_row"], ref["min_col"])
            return
        size = max(0, r2 - r1 + 1) * max(0, c2 - c1 + 1)
        populated = self.by_sheet[sheet]
        if size > len(populated):
            for (r, c) in populated:
                if r1 <= r <= r2 and c1 <= c <= c2:
                    yield (sheet, r, c)
        else:
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    if (r, c) in populated:
                        yield (sheet, r, c)

    def build_graph(self):
        for sheet, fcells in self.formula_cells.items():
            for cell in fcells.values():
                prec = set()
                for ref in cell.refs:
                    for k in self.expand(ref):
                        prec.add(k)
                        if len(prec) > MAX_EXPAND:
                            break
                self.precedents[cell.key] = prec

    # ---------------------------------------------------------------- checks
    def check_hardcoded_in_formula(self):
        for sheet, fcells in self.formula_cells.items():
            for cell in fcells.values():
                flagged, unit_only = [], True
                for num, ctx in cell.numbers:
                    if num in TRIVIAL_CONSTANTS:
                        continue
                    if any(fn in STRUCTURAL_ARG_FUNCS for fn in ctx[-1:]):
                        continue
                    flagged.append(num)
                    if num not in UNIT_CONSTANTS:
                        unit_only = False
                if not flagged:
                    continue
                nums = ", ".join(f"{n:g}" for n in flagged)
                if not cell.refs:
                    sev, msg = "medium", f"Formula is pure arithmetic on typed numbers ({nums})."
                elif unit_only:
                    sev, msg = "low", f"Unit-conversion constant(s) {nums} inside formula."
                else:
                    sev, msg = "medium", f"Hardcoded number(s) {nums} inside formula."
                self.add("hardcoded_in_formula", sev, sheet, cell.a1, msg, formula=cell.formula)

    def _runs(self, cells_by_line):
        """cells_by_line: {line: [(pos, Cell), ...]} -> list of contiguous runs."""
        runs = []
        for line, items in cells_by_line.items():
            items.sort(key=lambda x: x[0])
            run = [items[0]]
            for prev, cur in zip(items, items[1:]):
                if cur[0] == prev[0] + 1:
                    run.append(cur)
                else:
                    runs.append(run)
                    run = [cur]
            runs.append(run)
        return runs

    def check_consistency(self):
        seen = set()
        for sheet, fcells in self.formula_cells.items():
            rows, cols = defaultdict(list), defaultdict(list)
            for (r, c), cell in fcells.items():
                if cell.r1c1:
                    rows[r].append((c, cell))
                    cols[c].append((r, cell))
            for orient, groups in (("row", rows), ("column", cols)):
                for run in self._runs(groups):
                    for cells in self._segments([x[1] for x in run]):
                        self._flag_segment(sheet, orient, cells, seen)

    def _segments(self, cells):
        """Split a run where it switches between line formulas and aggregates, so a
        total row inside a column (or a total column in a row) is its own block."""
        segs, cur = [], []
        for c in cells:
            if cur and self._is_aggregate(c) != self._is_aggregate(cur[-1]):
                segs.append(cur)
                cur = []
            cur.append(c)
        if cur:
            segs.append(cur)
        return segs

    def _flag_segment(self, sheet, orient, cells, seen):
        # Drop a total that sums its own segment (row total at the end of a
        # total row): it summarizes the pattern rather than following it.
        keys = {c.key for c in cells}
        cells = [c for c in cells if not (self._is_aggregate(c) and
                 len(self.precedents.get(c.key, set()) & keys) >= 2)]
        if len(cells) < 4:
            return
        counts = Counter(c.r1c1 for c in cells)
        dom, n = counts.most_common(1)[0]
        if n < 3 or n / len(cells) < 0.7 or n == len(cells):
            return
        dom_cell = next(c for c in cells if c.r1c1 == dom)
        prev_ref = "C[-1]" if orient == "row" else "R[-1]"
        for i, c in enumerate(cells):
            if c.r1c1 == dom or c.key in seen:
                continue
            # In a block of totals, SUM vs AVERAGE is a design choice, not a slip.
            if self._is_aggregate(c) and c.funcs[0] != dom_cell.funcs[0]:
                continue
            # A seed cell ahead of a recursive pattern (=prior*(1+g)) is expected.
            if i == 0 and prev_ref in dom:
                continue
            seen.add(c.key)
            sev = "high" if 0 < i < len(cells) - 1 else ("low" if i == 0 else "medium")
            where = "first cell" if i == 0 else ("last cell" if i == len(cells) - 1
                                                 else "middle")
            self.add("inconsistent_formula", sev, sheet, c.a1,
                     f"Breaks the {orient} pattern ({where} of a {len(cells)}-cell run; "
                     f"{n} cells follow {dom_cell.a1}: {dom_cell.formula}).",
                     formula=c.formula)

    def _is_aggregate(self, cell):
        return bool(cell.funcs) and cell.funcs[0] in AGGREGATES and \
            cell.formula.upper().lstrip("=").startswith(cell.funcs[0])

    def check_pasted_over(self):
        for sheet, cells in self.by_sheet.items():
            fc = self.formula_cells[sheet]

            def pat(r, c):
                x = fc.get((r, c))
                return x.r1c1 if x else None

            for (r, c), cell in cells.items():
                if cell.formula or not is_number(cell.value):
                    continue
                left1, left2, right1 = pat(r, c - 1), pat(r, c - 2), pat(r, c + 1)
                up1, down1, up2 = pat(r - 1, c), pat(r + 1, c), pat(r - 2, c)
                if left1 and left1 == right1:
                    self.add("pasted_over_formula", "high", sheet, cell.a1,
                             f"Typed value {cell.value:g} between formulas that match "
                             f"({addr(c - 1, r)}: {fc[(r, c - 1)].formula}).", value=cell.value)
                elif left1 and left1 == left2 and right1 is None and not self.get(sheet, r, c + 1):
                    self.add("pasted_over_formula", "medium", sheet, cell.a1,
                             f"Typed value {cell.value:g} ends a row of formulas "
                             f"({addr(c - 1, r)}: {fc[(r, c - 1)].formula}).", value=cell.value)
                elif up1 and up1 == down1 and up1 == up2:
                    self.add("pasted_over_formula", "high", sheet, cell.a1,
                             f"Typed value {cell.value:g} inside a column of matching formulas "
                             f"({addr(c, r - 1)}: {fc[(r - 1, c)].formula}).", value=cell.value)

    def check_errors(self):
        err_cells = {}
        for key, cell in self.cells.items():
            v = cell.cached if cell.formula else cell.value
            if isinstance(v, str) and v.strip() in ERROR_VALUES:
                err_cells[key] = v.strip()
            elif cell.formula and "#REF!" in cell.formula.upper():
                err_cells[key] = "#REF!"
        roots = [k for k in err_cells
                 if not any(p in err_cells for p in self.precedents.get(k, ()))]
        propagated = len(err_cells) - len(roots)
        for k in sorted(roots):
            cell = self.cells[k]
            e = err_cells[k]
            sev = "critical" if e == "#REF!" else "high"
            downstream = sum(1 for j in err_cells if j != k and k in self.precedents.get(j, ()))
            self.add("error_value", sev, k[0], cell.a1,
                     f"{e} (root cause; {downstream} direct dependents also in error).",
                     formula=cell.formula, value=e)
        if propagated:
            self.add("error_value", "info", "*", "-",
                     f"{propagated} further cells show errors inherited from the roots above.")

    def check_refs_to_empty(self):
        for sheet, fcells in self.formula_cells.items():
            for cell in fcells.values():
                empties = []
                for ref in cell.refs:
                    if ref.get("kind") != "cell":
                        continue
                    if any(fn in BLANK_AWARE_FUNCS for fn in ref["context"]):
                        continue
                    if (ref["sheet"], ref["min_row"], ref["min_col"]) not in self.cells:
                        empties.append(ref["raw"])
                if empties:
                    self.add("ref_to_empty", "medium", sheet, cell.a1,
                             f"References empty cell(s): {', '.join(empties)}.", formula=cell.formula)

    def check_ranges(self):
        for sheet, fcells in self.formula_cells.items():
            for cell in fcells.values():
                for ref in cell.refs:
                    if ref.get("kind") in ("cols", "rows"):
                        includes_self = ref["sheet"] == sheet and (
                            (ref["kind"] == "cols" and ref["min_col"] <= cell.col <= ref["max_col"]) or
                            (ref["kind"] == "rows" and ref["min_row"] <= cell.row <= ref["max_row"]))
                        sev = "high" if includes_self else "medium"
                        self.add("whole_column_range", sev, sheet, cell.a1,
                                 f"Whole-{'column' if ref['kind'] == 'cols' else 'row'} range "
                                 f"{ref['raw']}" + (" includes this cell's own column/row."
                                                    if includes_self else "."),
                                 formula=cell.formula)
                if "SUM" not in cell.funcs:
                    continue
                for ref in cell.refs:
                    if ref.get("kind") not in ("range", "cols") or "SUM" not in ref["context"]:
                        continue
                    for k in self.expand(ref):
                        inner = self.formula_cells[k[0]].get((k[1], k[2]))
                        if not inner or inner.key == cell.key or not self._is_aggregate(inner):
                            continue
                        if inner.funcs[0] not in ("SUM", "SUBTOTAL"):
                            continue
                        inner_prec = self.precedents.get(inner.key, set())
                        covered = set(self.expand(ref))
                        if inner_prec and inner_prec & covered:
                            self.add("double_count", "high", sheet, cell.a1,
                                     f"Range {ref['raw']} includes subtotal {inner.a1} "
                                     f"({inner.formula}) and the rows it sums.",
                                     formula=cell.formula)
                            break

    def check_circular(self):
        graph = {k: [p for p in v if p in self.precedents] for k, v in self.precedents.items()}
        index, low, onstack, stack, sccs = {}, {}, set(), [], []
        counter = [0]
        for root in graph:
            if root in index:
                continue
            work = [(root, iter(graph[root]))]
            index[root] = low[root] = counter[0]; counter[0] += 1
            stack.append(root); onstack.add(root)
            while work:
                node, it = work[-1]
                advanced = False
                for nxt in it:
                    if nxt not in index:
                        index[nxt] = low[nxt] = counter[0]; counter[0] += 1
                        stack.append(nxt); onstack.add(nxt)
                        work.append((nxt, iter(graph.get(nxt, []))))
                        advanced = True
                        break
                    elif nxt in onstack:
                        low[node] = min(low[node], index[nxt])
                if advanced:
                    continue
                work.pop()
                if work:
                    low[work[-1][0]] = min(low[work[-1][0]], low[node])
                if low[node] == index[node]:
                    comp = []
                    while True:
                        w = stack.pop(); onstack.discard(w); comp.append(w)
                        if w == node:
                            break
                    if len(comp) > 1 or node in graph.get(node, []):
                        sccs.append(sorted(comp))
        for comp in sccs:
            first = comp[0]
            members = ", ".join(f"{s}!{addr(c, r)}" for s, r, c in comp[:12])
            more = f" (+{len(comp) - 12} more)" if len(comp) > 12 else ""
            self.add("circular_reference", "critical", first[0], addr(first[2], first[1]),
                     f"Circular loop of {len(comp)} cell(s): {members}{more}.",
                     formula=self.cells[first].formula, extra={"cells": [
                         f"{s}!{addr(c, r)}" for s, r, c in comp]})

    def _is_data(self, sheet, r, c):
        cell = self.get(sheet, r, c)
        if cell is None:
            return False
        if cell.formula:
            if self._is_aggregate(cell):
                return False
            return cell.cached is None or is_number(cell.cached)
        return is_number(cell.value)

    def check_sum_omission(self):
        for sheet, fcells in self.formula_cells.items():
            for cell in fcells.values():
                if not self._is_aggregate(cell) or cell.funcs[0] != "SUM":
                    continue
                refs = [r for r in cell.refs if r.get("sheet") == sheet and
                        r.get("kind") in ("range", "cell")]
                if not refs:
                    continue
                vertical = all(r["min_col"] == r["max_col"] == cell.col for r in refs)
                horizontal = all(r["min_row"] == r["max_row"] == cell.row for r in refs)
                if not (vertical or horizontal):
                    continue
                covered = set()
                for r in refs:
                    if vertical:
                        covered.update(range(r["min_row"], r["max_row"] + 1))
                    else:
                        covered.update(range(r["min_col"], r["max_col"] + 1))
                lo, hi = min(covered), max(covered)
                me = cell.row if vertical else cell.col

                def data(p):
                    return self._is_data(sheet, p, cell.col) if vertical else \
                        self._is_data(sheet, cell.row, p)

                missed = []
                # gaps between the formula and the range, and holes inside a split range
                span = range(hi + 1, me) if me > hi else range(me + 1, lo)
                missed += [p for p in span if data(p)]
                missed += [p for p in range(lo, hi + 1) if p not in covered and data(p)]
                # contiguous data just beyond the far edge of the range
                edge = lo - 1 if me > hi else hi + 1
                member_pats = set()
                for p in covered:
                    x = self.get(sheet, p, cell.col) if vertical else self.get(sheet, cell.row, p)
                    member_pats.add(x.r1c1 if x and x.formula else "const")

                def sibling(p):
                    # Only a line item that looks like the summed ones counts as omitted;
                    # a different calculation directly above (e.g. gross profit) does not.
                    x = self.get(sheet, p, cell.col) if vertical else self.get(sheet, cell.row, p)
                    return (x.r1c1 if x.formula else "const") in member_pats

                while edge > 0 and data(edge) and sibling(edge) and edge not in covered \
                        and edge != me:
                    lbl = self.label_for_row(sheet, edge) if vertical else ""
                    if TOTAL_WORDS.search(lbl):
                        break
                    missed.append(edge)
                    edge += -1 if me > hi else 1
                if missed:
                    names = [addr(cell.col, p) if vertical else addr(p, cell.row)
                             for p in sorted(set(missed))]
                    self.add("sum_range_omission", "high", sheet, cell.a1,
                             f"SUM skips adjacent data cell(s): {', '.join(names)}.",
                             formula=cell.formula)

    def check_footing(self):
        # typed totals vs their detail
        for sheet, cells in self.by_sheet.items():
            for (r, c), cell in cells.items():
                if cell.formula or not is_number(cell.value):
                    continue
                row_label = self.label_for_row(sheet, r)
                col_label = self.label_for_col(sheet, r, c)
                for orient, label in (("row", row_label), ("col", col_label)):
                    if not TOTAL_WORDS.search(label or ""):
                        continue
                    parts, unknown = [], False
                    p = (r - 1) if orient == "row" else (c - 1)
                    while p > 0:
                        x = self.get(sheet, p, c) if orient == "row" else self.get(sheet, r, p)
                        if x is None:
                            break
                        v = self.numeric_value(x)
                        if v is None:
                            if x.formula and x.cached is None:
                                unknown = True
                            break
                        parts.append(v)
                        p -= 1
                    if len(parts) < 2:
                        continue
                    if unknown:
                        self.add("hardcoded_total", "medium", sheet, cell.a1,
                                 f"Typed total {cell.value:g} under '{label}' (detail not "
                                 f"calculated, footing not verified).", value=cell.value)
                        break
                    s = sum(parts)
                    if abs(s - cell.value) > max(0.01, 1e-6 * abs(s)):
                        self.add("total_does_not_foot", "high", sheet, cell.a1,
                                 f"Typed total {cell.value:g} but the {len(parts)} cells "
                                 f"{'above' if orient == 'row' else 'to the left'} sum to {s:g} "
                                 f"(off by {cell.value - s:g}).", value=cell.value)
                    else:
                        self.add("hardcoded_total", "medium", sheet, cell.a1,
                                 f"Typed total {cell.value:g} foots today but will not update.",
                                 value=cell.value)
                    break
        if not self.has_cache:
            return
        # calculated SUMs whose cached value disagrees with cached inputs
        for sheet, fcells in self.formula_cells.items():
            for cell in fcells.values():
                if not is_number(cell.cached) or not cell.formula.upper().startswith("=SUM(") \
                        or cell.formula.count("(") != 1:
                    continue
                total, ok = 0.0, True
                for ref in cell.refs:
                    if ref.get("kind") not in ("cell", "range"):
                        ok = False
                        break
                    for k in self.expand(ref):
                        x = self.cells.get(k)
                        v = x.display_value() if x else None
                        if x and x.formula and v is None:
                            ok = False
                            break
                        if is_number(v):
                            total += v
                if ok and abs(total - cell.cached) > max(0.01, 1e-6 * abs(total)):
                    self.add("stale_cached_value", "high", sheet, cell.a1,
                             f"Stored value {cell.cached:g} but its inputs sum to {total:g}.",
                             formula=cell.formula, value=cell.cached)

    def check_links_and_volatile(self):
        inbound = Counter()
        for sheet, fcells in self.formula_cells.items():
            for cell in fcells.values():
                ext = [r["raw"] for r in cell.refs if r.get("kind") == "external"]
                ext += sorted({m.upper() for m in SHEETS_LIVE.findall(cell.formula)})
                if ext:
                    self.add("external_link", "high", sheet, cell.a1,
                             f"Links to another workbook: {', '.join(ext)}.", formula=cell.formula)
                for r in cell.refs:
                    if r.get("kind") not in ("external", "unknown") and r["sheet"] != sheet:
                        inbound[(sheet, r["sheet"])] += 1
                vol = sorted(set(cell.funcs) & VOLATILE)
                if vol:
                    sev = "medium" if set(vol) & {"INDIRECT", "OFFSET", "RAND", "RANDBETWEEN",
                                                  "RANDARRAY"} else "low"
                    self.add("volatile_function", sev, sheet, cell.a1,
                             f"Uses volatile function(s): {', '.join(vol)}.", formula=cell.formula)
        if self.ext_link_count and not any(f["check"] == "external_link" for f in self.findings):
            self.add("external_link", "medium", "*", "-",
                     f"Workbook stores {self.ext_link_count} external link definition(s) "
                     f"(Data > Edit Links) even though no formula uses them.")
        for (src, dst), n in sorted(inbound.items()):
            self.add("cross_sheet_link", "info", src, "-",
                     f"{n} formula reference(s) from '{src}' into '{dst}'.")

    def check_hidden(self):
        referenced = set()
        for prec in self.precedents.values():
            referenced |= prec
        for sheet, state in self.sheet_state.items():
            if state == "visible":
                continue
            nf = len(self.formula_cells[sheet])
            nv = len(self.by_sheet[sheet])
            used = sum(1 for k in referenced if k[0] == sheet)
            sev = "high" if state == "veryHidden" else ("medium" if nf or used else "low")
            self.add("hidden_content", sev, sheet, "-",
                     f"Sheet is {state} with {nv} populated cells ({nf} formulas, {used} "
                     f"referenced by other formulas).")
        for sheet in self.sheetnames:
            for kind, hidden in (("row", self.hidden_rows[sheet]), ("column", self.hidden_cols[sheet])):
                for h in sorted(hidden):
                    content = [x for (r, c), x in self.by_sheet[sheet].items()
                               if (r if kind == "row" else c) == h]
                    if not content:
                        continue
                    nf = sum(1 for x in content if x.formula)
                    used = sum(1 for x in content if x.key in referenced)
                    sev = "medium" if nf or used else "low"
                    label = h if kind == "row" else get_column_letter(h)
                    self.add("hidden_content", sev, sheet, f"{kind} {label}",
                             f"Hidden {kind} {label} holds {len(content)} cells ({nf} formulas, "
                             f"{used} feeding visible formulas).")

    def check_units_and_signs(self):
        for sheet, cells in self.by_sheet.items():
            rows = defaultdict(list)
            for (r, c), cell in cells.items():
                if cell.formula or not is_number(cell.value):
                    continue
                if "%" in cell.number_format and 1 < abs(cell.value) <= 100:
                    self.add("percent_as_whole_number", "medium", sheet, cell.a1,
                             f"Percent-formatted input holds {cell.value:g} "
                             f"(displays as {cell.value * 100:g}%).", value=cell.value)
                rows[r].append(cell)
            expense_signs = {}
            for r, items in rows.items():
                vals = [x.value for x in items if x.value != 0]
                if len(vals) >= 4:
                    mags = sorted(abs(v) for v in vals)
                    med = mags[len(mags) // 2]
                    for x in items:
                        if med and x.value and (abs(x.value) / med >= 90 or abs(x.value) / med <= 1 / 90):
                            self.add("unit_outlier", "low", sheet, x.a1,
                                     f"Input {x.value:g} is ~{abs(x.value) / med:.3g}x the row "
                                     f"median {med:g}.", value=x.value)
                label = self.label_for_row(sheet, r)
                if EXPENSE_WORDS.search(label) and vals:
                    pos, neg = sum(v > 0 for v in vals), sum(v < 0 for v in vals)
                    if pos and neg:
                        self.add("mixed_sign_convention", "low", sheet, f"row {r}",
                                 f"'{label}' mixes positive ({pos}) and negative ({neg}) inputs.")
                    elif pos or neg:
                        expense_signs[r] = (label, "+" if pos else "-")
            signs = {s for _, s in expense_signs.values()}
            if len(signs) > 1:
                pos_rows = [f"{r} ({l})" for r, (l, s) in expense_signs.items() if s == "+"][:4]
                neg_rows = [f"{r} ({l})" for r, (l, s) in expense_signs.items() if s == "-"][:4]
                self.add("mixed_sign_convention", "low", sheet, "-",
                         f"Expense rows use both conventions. Positive: {'; '.join(pos_rows)}. "
                         f"Negative: {'; '.join(neg_rows)}.")

    # ---------------------------------------------------------------- driver
    def run(self):
        self.load()
        self.build_graph()
        nformulas = sum(len(v) for v in self.formula_cells.values())
        if nformulas and not self.has_cache:
            self.add("no_cached_values", "info", "*", "-",
                     f"{nformulas} formulas have no stored results.")
        self.check_errors()
        self.check_circular()
        self.check_hardcoded_in_formula()
        self.check_pasted_over()
        self.check_consistency()
        self.check_refs_to_empty()
        self.check_ranges()
        self.check_sum_omission()
        self.check_footing()
        self.check_links_and_volatile()
        self.check_hidden()
        self.check_units_and_signs()
        self.findings.sort(key=lambda f: (SEVERITY_ORDER.index(f["severity"]), f["check"],
                                          f["sheet"], _cell_sort(f["cell"])))
        for i, f in enumerate(self.findings, 1):
            f["id"] = f"F{i:03d}"
        return {
            "file": os.path.basename(self.path),
            "sheets": self.sheetnames,
            "formula_count": nformulas,
            "cell_count": len(self.cells),
            "cached_values": self.has_cache,
            "summary": dict(Counter(f["severity"] for f in self.findings)),
            "by_check": dict(Counter(f["check"] for f in self.findings)),
            "findings": self.findings,
        }


def _cell_sort(a1):
    m = re.match(r"([A-Z]+)(\d+)$", a1 or "")
    if not m:
        return (0, 0, a1 or "")
    return (int(m.group(2)), column_index_from_string(m.group(1)), "")


# --------------------------------------------------------------------------- output

def to_markdown(report, max_per_check=25):
    lines = [f"# Spreadsheet audit: {report['file']}", ""]
    s = report["summary"]
    lines.append(f"Sheets: {', '.join(report['sheets'])}  ")
    lines.append(f"Populated cells: {report['cell_count']}, formulas: {report['formula_count']}, "
                 f"cached values: {'yes' if report['cached_values'] else 'NO (value checks skipped)'}")
    lines.append("")
    lines.append("| Severity | Count |")
    lines.append("|---|---|")
    for sev in SEVERITY_ORDER:
        lines.append(f"| {sev} | {s.get(sev, 0)} |")
    lines.append("")
    for sev in SEVERITY_ORDER:
        group = [f for f in report["findings"] if f["severity"] == sev]
        if not group:
            continue
        lines.append(f"## {sev.capitalize()}")
        lines.append("")
        by_check = defaultdict(list)
        for f in group:
            by_check[f["check"]].append(f)
        for check, items in by_check.items():
            lines.append(f"### {check.replace('_', ' ')} ({len(items)})")
            lines.append("")
            lines.append(f"Why it matters: {items[0]['why']}  ")
            lines.append(f"Fix: {items[0]['fix']}")
            lines.append("")
            for f in items[:max_per_check]:
                loc = f"{f['sheet']}!{f['cell']}" if f["cell"] != "-" else f["sheet"]
                formula = f"  `{f['formula']}`" if f.get("formula") else ""
                lines.append(f"- **{f['id']}** `{loc}` {f['message']}{formula}")
            if len(items) > max_per_check:
                lines.append(f"- ... and {len(items) - max_per_check} more (see JSON).")
            lines.append("")
    if not report["findings"]:
        lines.append("No structural issues found. Review assumptions manually (see SKILL.md).")
    return "\n".join(lines)


def recalc_copy(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("warning: --recalc requested but LibreOffice (soffice) not found; "
              "auditing without recalculation", file=sys.stderr)
        return path
    outdir = tempfile.mkdtemp(prefix="xlsx-recalc-")
    src = os.path.join(outdir, "src_" + os.path.basename(path))
    shutil.copy(path, src)
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", os.path.join(outdir, "out"), src],
                   check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    out = os.path.join(outdir, "out", os.path.splitext(os.path.basename(src))[0] + ".xlsx")
    if not os.path.exists(out):
        print("warning: LibreOffice conversion failed; auditing original", file=sys.stderr)
        return path
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path")
    ap.add_argument("--json", help="write JSON report here")
    ap.add_argument("--md", help="write markdown report here")
    ap.add_argument("--recalc", action="store_true", help="recalculate via LibreOffice first")
    ap.add_argument("--max-per-check", type=int, default=25)
    args = ap.parse_args(argv)
    if not os.path.exists(args.path):
        sys.exit(f"not found: {args.path}")
    if args.path.lower().endswith(".xls"):
        sys.exit("legacy .xls is not supported by openpyxl; save as .xlsx first "
                 "(or: soffice --headless --convert-to xlsx file.xls)")
    path = recalc_copy(args.path) if args.recalc else args.path
    report = Auditor(path, args.max_per_check).run()
    report["file"] = os.path.basename(args.path)
    report["recalculated"] = path != args.path
    md = to_markdown(report, args.max_per_check)
    if args.json:
        with open(args.json, "w") as fh:
            json.dump(report, fh, indent=2, default=str)
    if args.md:
        with open(args.md, "w") as fh:
            fh.write(md)
    if not args.json and not args.md:
        print(md)
    else:
        print(json.dumps({"file": report["file"], "summary": report["summary"],
                          "by_check": report["by_check"]}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
