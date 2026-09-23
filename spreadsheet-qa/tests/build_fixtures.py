#!/usr/bin/env python3
"""Write the synthetic fixtures used by run_tests.py.

Planted traps:
  orders.csv     - a "Grand Total" row at the bottom that sums every amount
                   (across currencies); amounts stored as text ("1,200.00",
                   "$350.00"); USD/EUR/GBP mixed in one amount column; UTC
                   timestamps where the UTC day differs from the New York day;
                   a cancelled order; orders just outside Q3 on each side.
  customers.csv  - customer C003 appears twice (segment changed from
                   Mid-Market to Enterprise), so a naive join fans out.
  refunds.csv    - one refund dated after the quarter for an in-quarter order.
  fx_rates.csv   - fixed rates to USD.
  pipeline.xlsx  - two title rows above the header, Excel serial dates, a
                   hidden row, an autofilter, a merged title, and a Total row.
"""
import csv
import os

ORDERS = [
    # order_id, customer_id, created_at (UTC), currency, amount, status
    ("O1001", "C001", "2026-06-30T23:30:00Z", "USD", "500.00", "paid"),
    ("O1002", "C001", "2026-07-03T15:00:00Z", "USD", "1,200.00", "paid"),
    ("O1003", "C002", "2026-07-10T09:00:00Z", "USD", "300.00", "paid"),
    ("O1004", "C003", "2026-07-15T12:00:00Z", "EUR", "1000.00", "paid"),
    ("O1005", "C004", "2026-07-20T18:00:00Z", "GBP", "400.00", "paid"),
    ("O1006", "C005", "2026-07-31T23:00:00Z", "USD", "2000.00", "paid"),
    ("O1007", "C002", "2026-08-01T02:30:00Z", "USD", "150.00", "paid"),
    ("O1008", "C003", "2026-08-05T10:00:00Z", "EUR", "500.00", "paid"),
    ("O1009", "C001", "2026-08-12T14:00:00Z", "USD", "$350.00", "paid"),
    ("O1010", "C004", "2026-08-20T16:00:00Z", "GBP", "200.00", "cancelled"),
    ("O1011", "C005", "2026-08-31T23:30:00Z", "USD", "800.00", "paid"),
    ("O1012", "C002", "2026-09-01T03:00:00Z", "USD", "250.00", "paid"),
    ("O1013", "C003", "2026-09-10T11:00:00Z", "EUR", "700.00", "paid"),
    ("O1014", "C001", "2026-09-15T13:00:00Z", "USD", "450.00", "paid"),
    ("O1015", "C005", "2026-09-30T22:00:00Z", "USD", "600.00", "paid"),
    ("O1016", "C004", "2026-10-01T01:00:00Z", "GBP", "300.00", "paid"),
]
GRAND_TOTAL = 500 + 1200 + 300 + 1000 + 400 + 2000 + 150 + 500 + 350 + 200 + 800 + 250 + 700 + 450 + 600 + 300

CUSTOMERS = [
    ("customer_id", "name", "segment", "country", "updated_at"),
    ("C001", "Acme Tools", "SMB", "US", "2026-01-05"),
    ("C002", "Birch Dental", "SMB", "US", "2026-01-09"),
    ("C003", "Cobalt Logistics", "Mid-Market", "DE", "2026-02-01"),
    ("C003", "Cobalt Logistics GmbH", "Enterprise", "DE", "2026-06-15"),
    ("C004", "Dune Coffee", "SMB", "GB", "2026-03-02"),
    ("C005", "Elm Health", "Enterprise", "US", "2026-04-11"),
    ("C006", "Fjord Studio", "Mid-Market", "GB", "2026-05-20"),
]

REFUNDS = [
    ("refund_id", "order_id", "refunded_at", "currency", "amount"),
    ("R01", "O1002", "2026-07-20T10:00:00Z", "USD", "200.00"),
    ("R02", "O1004", "2026-08-01T10:00:00Z", "EUR", "100.00"),
    ("R03", "O1014", "2026-10-02T10:00:00Z", "USD", "450.00"),
]

FX = [("currency", "usd_per_unit"), ("USD", "1.0"), ("EUR", "1.10"), ("GBP", "1.25")]


def _write(path, rows):
    with open(path, "w", newline="") as f:
        csv.writer(f).writerows(rows)


def build(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    _write(os.path.join(out_dir, "orders.csv"),
           [("order_id", "customer_id", "created_at", "currency", "amount", "status")] + ORDERS
           + [("", "Grand Total", "", "", f"{GRAND_TOTAL:,.2f}", "")])
    _write(os.path.join(out_dir, "customers.csv"), CUSTOMERS)
    _write(os.path.join(out_dir, "refunds.csv"), REFUNDS)
    _write(os.path.join(out_dir, "fx_rates.csv"), FX)
    build_xlsx(os.path.join(out_dir, "pipeline.xlsx"))
    return out_dir


def build_xlsx(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    ws["A1"] = "Q3 Pipeline Report"
    ws.merge_cells("A1:D1")
    ws["A2"] = "Exported 2026-09-30 by sales ops"
    ws.append(["deal_id", "owner", "close_date", "amount"])
    deals = [("D-1", "Ana", 46204, 12000), ("D-2", "Ben", 46215, 8000), ("D-3", "Ana", 46230, 5000),
             ("D-4", "Cy", 46241, 7000), ("D-5", "Ben", 46260, 3000)]
    for d in deals:
        ws.append(list(d))
    ws.append(["Total", None, None, sum(d[3] for d in deals)])
    ws.row_dimensions[6].hidden = True  # D-3 hidden
    ws.auto_filter.ref = "A3:D8"
    wb.save(path)


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1] if len(sys.argv) > 1 else "fixtures"))
